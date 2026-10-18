from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import date, datetime, timezone
from typing import Annotated
from urllib.parse import quote
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.core.spreadsheet import neutralise_formulas
from app.core.config import get_settings
from app.core.database import get_db
from app.core.models import (
    AuditLog, Equipment, EquipmentDocument, Issue, OMTask, Plant, User,
    UserEquipmentTypeAssignment,
)
from app.core.security import (
    get_current_user, require_csrf, require_permission, user_equipment_matches_record,
    user_scoped_equipment_ids, user_scoped_plant_ids,
)
from app.plants.repository import EquipmentRepository, PlantRepository
from app.plants.schemas import (
    EquipmentBulkCreate, EquipmentCreate, EquipmentMove, EquipmentOut, EquipmentPage,
    EquipmentUpdate, PlantCreate, PlantOut, PlantUpdate,
)
from app.plants.service import EquipmentService, PlantService
from app.storage.service import SupabaseStorage

router=APIRouter(prefix="/api", tags=["Plants"])
settings=get_settings()


@router.get("/plants", response_model=list[PlantOut])
def list_plants(db:Annotated[Session,Depends(get_db)], user:Annotated[User,Depends(get_current_user)], q:str|None=None,district:str|None=None,project:str|None=None,engineer:str|None=None,offset:int=0,limit:int=Query(50,le=200)):
    return PlantRepository(db).list(q=q,district=district,project=project,engineer=engineer,allowed_plant_ids=user_scoped_plant_ids(db,user),offset=offset,limit=limit)


@router.post("/plants", response_model=PlantOut, status_code=201, dependencies=[Depends(require_csrf)])
def create_plant(data:PlantCreate,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]): return PlantService(db).create(data,user)


@router.patch("/plants/{plant_id}", response_model=PlantOut, dependencies=[Depends(require_csrf)])
def update_plant(plant_id:UUID,data:PlantUpdate,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]): return PlantService(db).update(plant_id,data)


@router.delete("/plants/{plant_id}", status_code=204, dependencies=[Depends(require_csrf)])
def delete_plant(plant_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]): PlantService(db).delete(plant_id); return Response(status_code=204)


def _equipment_scope(db: Session, user: User) -> dict:
    return {
        "allowed_plant_ids": user_scoped_plant_ids(db,user) if user.role != "Assignee" else None,
        "allowed_equipment_ids": user_scoped_equipment_ids(db,user) if user.role == "Assignee" else None,
    }


def _ensure_equipment_read(db: Session, user: User, row: Equipment) -> None:
    if not user_equipment_matches_record(db,user,row.plant_id,row.id,row.equipment_code or row.equipment_name):
        raise HTTPException(403,"This equipment is outside your assigned scope")


def _type_owners(db: Session, equipment_type: str | None) -> list[str]:
    if not equipment_type:
        return []
    rows=db.execute(
        select(User.display_name)
        .join(UserEquipmentTypeAssignment,UserEquipmentTypeAssignment.user_id==User.id)
        .where(User.is_active.is_(True),func.lower(UserEquipmentTypeAssignment.equipment_type)==equipment_type.strip().lower())
        .order_by(User.display_name)
    ).scalars().all()
    return list(rows)


def _equipment_out(
    db: Session,
    row: Equipment,
    *,
    in_charge_name: str | None = None,
    type_owners: list[str] | None = None,
    prefetched: bool = False,
) -> dict:
    if prefetched:
        person_name=in_charge_name
        owners=type_owners or []
    else:
        person=db.get(User,row.in_charge_user_id) if row.in_charge_user_id else None
        person_name=person.display_name if person else None
        owners=_type_owners(db,row.equipment_type)
    return {
        "id":row.id,"plant_id":row.plant_id,"equipment_code":row.equipment_code,"equipment_name":row.equipment_name,
        "equipment_type":row.equipment_type,"maker":row.maker,"model":row.model,"serial_number":row.serial_number,
        "status":row.status,"discom":row.discom,"feeder_or_gss":row.feeder_or_gss,"rating":row.rating,
        "capacity_kw":float(row.capacity_kw) if row.capacity_kw is not None else None,"quantity":row.quantity or 1,
        "location_in_plant":row.location_in_plant,"manufacture_year":row.manufacture_year,"commissioned_on":row.commissioned_on,
        "warranty_start":row.warranty_start,"warranty_end":row.warranty_end,"amc_vendor":row.amc_vendor,"amc_end":row.amc_end,
        "supplier":row.supplier,"purchase_order_no":row.purchase_order_no,"criticality":row.criticality,"notes":row.notes,
        "in_charge_user_id":row.in_charge_user_id,"in_charge_name":person_name,
        "last_service_on":row.last_service_on,"next_service_due":row.next_service_due,"metadata":row.metadata_json or {},
        "type_owners":owners,"created_at":row.created_at,"updated_at":row.updated_at,
    }


def _equipment_out_many(db: Session, rows: list[Equipment]) -> list[dict]:
    if not rows:
        return []

    user_ids={row.in_charge_user_id for row in rows if row.in_charge_user_id}
    people={
        person.id:person.display_name
        for person in db.scalars(select(User).where(User.id.in_(user_ids)))
    } if user_ids else {}
    equipment_types={row.equipment_type.strip().casefold() for row in rows if row.equipment_type and row.equipment_type.strip()}
    owners_by_type:dict[str,list[str]]={}
    if equipment_types:
        assignments=db.execute(
            select(UserEquipmentTypeAssignment.equipment_type,User.display_name)
            .join(User,User.id==UserEquipmentTypeAssignment.user_id)
            .where(
                User.is_active.is_(True),
                func.lower(UserEquipmentTypeAssignment.equipment_type).in_(equipment_types),
            )
            .order_by(User.display_name)
        ).all()
        for equipment_type,display_name in assignments:
            owners_by_type.setdefault(equipment_type.strip().casefold(),[]).append(display_name)

    return [
        _equipment_out(
            db,row,
            in_charge_name=people.get(row.in_charge_user_id) if row.in_charge_user_id else None,
            type_owners=owners_by_type.get((row.equipment_type or "").strip().casefold(),[]),
            prefetched=True,
        )
        for row in rows
    ]


def _list_filters(plant_id,q,equipment_type,maker,status,discom,in_charge_user_id,warranty,no_in_charge,maker_missing,serial_missing=False):
    return dict(plant_id=plant_id,q=q,equipment_type=equipment_type,maker=maker,status=status,discom=discom,
                in_charge_user_id=in_charge_user_id,warranty=warranty,no_in_charge=no_in_charge,maker_missing=maker_missing,serial_missing=serial_missing)


@router.get("/equipment", response_model=list[EquipmentOut])
def list_equipment(
    db:Annotated[Session,Depends(get_db)], user:Annotated[User,Depends(get_current_user)], plant_id:UUID|None=None,
    q:str|None=None, equipment_type:str|None=None, maker:str|None=None, status:str|None=None, discom:str|None=None,
    in_charge_user_id:UUID|None=None, warranty:str|None=None, no_in_charge:bool=False, maker_missing:bool=False, serial_missing:bool=False,
    offset:int=0,limit:int=Query(100,le=300),
):
    filters=_list_filters(plant_id,q,equipment_type,maker,status,discom,in_charge_user_id,warranty,no_in_charge,maker_missing,serial_missing)
    rows=EquipmentRepository(db).list(**filters,**_equipment_scope(db,user),offset=offset,limit=limit)
    return _equipment_out_many(db,rows)


@router.get("/equipment/page", response_model=EquipmentPage)
def page_equipment(
    db:Annotated[Session,Depends(get_db)], user:Annotated[User,Depends(get_current_user)], plant_id:UUID|None=None,
    q:str|None=None, equipment_type:str|None=None, maker:str|None=None, status:str|None=None, discom:str|None=None,
    in_charge_user_id:UUID|None=None, warranty:str|None=None, no_in_charge:bool=False, maker_missing:bool=False, serial_missing:bool=False,
    offset:int=0,limit:int=Query(50,ge=10,le=200),
):
    filters=_list_filters(plant_id,q,equipment_type,maker,status,discom,in_charge_user_id,warranty,no_in_charge,maker_missing,serial_missing)
    scoped={**filters,**_equipment_scope(db,user)}
    repo=EquipmentRepository(db)
    rows=repo.list(**scoped,offset=offset,limit=limit)
    return {"items":_equipment_out_many(db,rows),"total":repo.count(**scoped),"offset":offset,"limit":limit}


@router.post("/equipment", response_model=EquipmentOut, status_code=201, dependencies=[Depends(require_csrf)])
def create_equipment(data:EquipmentCreate,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    return _equipment_out(db,EquipmentService(db).create(data,user))


@router.post("/equipment/bulk", response_model=list[EquipmentOut], status_code=201, dependencies=[Depends(require_csrf)])
def create_equipment_bulk(data:EquipmentBulkCreate,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    return _equipment_out_many(db,EquipmentService(db).bulk_create(data,user))


@router.patch("/equipment/{equipment_id}", response_model=EquipmentOut, dependencies=[Depends(require_csrf)])
def update_equipment(equipment_id:UUID,data:EquipmentUpdate,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    return _equipment_out(db,EquipmentService(db).update(equipment_id,data,user))


@router.post("/equipment/{equipment_id}/move", response_model=EquipmentOut, dependencies=[Depends(require_csrf)])
def move_equipment(equipment_id:UUID,data:EquipmentMove,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    return _equipment_out(db,EquipmentService(db).move(equipment_id,data.plant_id,user))


@router.post("/equipment/{equipment_id}/retire", response_model=EquipmentOut, dependencies=[Depends(require_csrf)])
def retire_equipment(equipment_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    return _equipment_out(db,EquipmentService(db).retire(equipment_id,user))


@router.delete("/equipment/{equipment_id}", status_code=204, dependencies=[Depends(require_csrf)])
def delete_equipment(equipment_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    EquipmentService(db).delete(equipment_id,user); return Response(status_code=204)


@router.get("/equipment/{equipment_id}/detail")
def equipment_detail(equipment_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(get_current_user)]):
    row=db.get(Equipment,equipment_id)
    if not row: raise HTTPException(404,"Equipment not found")
    _ensure_equipment_read(db,user,row)
    refs={x.strip().casefold() for x in [row.equipment_code or "",row.equipment_name or ""] if x.strip()}
    task_stmt=select(OMTask).where(OMTask.plant_id==row.plant_id)
    issue_stmt=select(Issue).where(Issue.plant_id==row.plant_id,Issue.deleted_at.is_(None))
    tasks=[t for t in db.scalars(task_stmt.order_by(OMTask.due_on.desc())) if (t.equipment_ref or "").strip().casefold() in refs]
    issues=[i for i in db.scalars(issue_stmt.order_by(Issue.created_at.desc())) if i.equipment_id==row.id or (i.equipment_ref or "").strip().casefold() in refs]
    history=list(db.scalars(select(AuditLog).where(AuditLog.entity_type=="equipment",AuditLog.entity_id==str(row.id)).order_by(AuditLog.created_at.desc())))
    docs=list(db.scalars(select(EquipmentDocument).where(EquipmentDocument.equipment_id==row.id).order_by(EquipmentDocument.created_at.desc())))
    open_issues=[i for i in issues if i.status not in {"Resolved","Closed"}]
    repair_hours=[]
    for i in issues:
        closed=(i.closure or {}).get("closed_at") if isinstance(i.closure,dict) else None
        if closed and i.created_at:
            try: repair_hours.append(max(0,(datetime.fromisoformat(str(closed).replace("Z","+00:00"))-i.created_at).total_seconds()/3600))
            except Exception: pass
    return {
        "equipment":_equipment_out(db,row),
        "assignment_history":[{"at":a.created_at.isoformat() if a.created_at else None,"by":a.actor_name,"action":a.action,"before":a.before_state,"after":a.after_state,"metadata":a.metadata_json or {}} for a in history if (a.action or "").lower().find("move")>=0 or (a.before_state or {}).get("in_charge_user_id")!=(a.after_state or {}).get("in_charge_user_id")],
        "maintenance":[{"id":t.work_order_no,"name":t.task_name,"status":t.status,"due_on":t.due_on.isoformat() if t.due_on else None,"actual_completed_on":t.actual_completed_on.isoformat() if t.actual_completed_on else None} for t in tasks[:100]],
        "issues":[{"id":str(i.id),"ticket_no":i.ticket_no,"status":i.status,"severity":"Critical" if i.severity=="High" else i.severity,"created_at":i.created_at.isoformat() if i.created_at else None} for i in issues[:100]],
        "issue_metrics":{"open_count":len(open_issues),"mttr_hours":round(sum(repair_hours)/len(repair_hours),1) if repair_hours else None,"mtbf_hours":None},
        "documents":[{"id":str(d.id),"kind":d.kind,"name":d.original_name,"content_type":d.content_type,"size":d.size_bytes,"created_at":d.created_at.isoformat() if d.created_at else None,"url":f"/api/equipment/{row.id}/documents/{d.id}"} for d in docs],
    }


_DOC_KINDS={"datasheet","warranty_certificate","test_report"}
_DOC_TYPES={"application/pdf":"pdf","image/jpeg":"jpg","image/png":"png","image/webp":"webp","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":"xlsx"}


def _safe_document_filename(filename: str | None, fallback: str = "document") -> str:
    basename=(filename or "").replace("\\","/").rsplit("/",1)[-1]
    basename="".join(character for character in basename if character.isprintable()).strip(" .")
    return basename[:255] or fallback


def _document_content_disposition(filename: str) -> str:
    safe_name=_safe_document_filename(filename)
    ascii_name=safe_name.encode("ascii","ignore").decode("ascii")
    fallback=re.sub(r"[^A-Za-z0-9._ -]","_",ascii_name).strip(" .") or "document"
    return f"inline; filename=\"{fallback}\"; filename*=UTF-8''{quote(safe_name,safe='')}"


@router.post("/equipment/{equipment_id}/documents", dependencies=[Depends(require_csrf)])
def upload_equipment_document(
    equipment_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))],
    kind:str=Form(...),file:UploadFile=File(...),
):
    row=db.get(Equipment,equipment_id)
    if not row: raise HTTPException(404,"Equipment not found")
    if kind not in _DOC_KINDS: raise HTTPException(422,"Document kind must be datasheet, warranty_certificate or test_report")
    content_type=(file.content_type or "").lower()
    if content_type not in _DOC_TYPES: raise HTTPException(415,"Supported files: PDF, JPEG, PNG, WebP or XLSX")
    data=file.file.read(settings.max_upload_mb*1024*1024+1)
    if len(data)>settings.max_upload_mb*1024*1024: raise HTTPException(413,f"File exceeds {settings.max_upload_mb} MB")
    if len(data)<16: raise HTTPException(422,"File is empty or invalid")
    path=f"equipment/{row.id}/{kind}-{uuid.uuid4().hex}.{_DOC_TYPES[content_type]}"
    SupabaseStorage().upload(path,data,content_type)
    original_name=_safe_document_filename(file.filename,f"{kind}.{_DOC_TYPES[content_type]}")
    doc=EquipmentDocument(equipment_id=row.id,kind=kind,original_name=original_name,storage_path=path,content_type=content_type,size_bytes=len(data),uploaded_by=user.id)
    db.add(doc); db.flush()
    db.add(AuditLog(user_id=user.id,actor_name=user.display_name,action="Uploaded equipment document",module="Equipment Register",entity_type="equipment",entity_id=str(row.id),metadata_json={"document_id":str(doc.id),"kind":kind,"name":doc.original_name}))
    db.commit(); db.refresh(doc)
    return {"id":str(doc.id),"kind":doc.kind,"name":doc.original_name,"size":doc.size_bytes,"url":f"/api/equipment/{row.id}/documents/{doc.id}"}


@router.get("/equipment/{equipment_id}/documents/{document_id}")
def get_equipment_document(equipment_id:UUID,document_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(get_current_user)]):
    row=db.get(Equipment,equipment_id)
    if not row: raise HTTPException(404,"Equipment not found")
    _ensure_equipment_read(db,user,row)
    doc=db.get(EquipmentDocument,document_id)
    if not doc or doc.equipment_id!=row.id: raise HTTPException(404,"Document not found")
    data,content_type=SupabaseStorage().download(doc.storage_path)
    return Response(content=data,media_type=content_type,headers={"Content-Disposition":_document_content_disposition(doc.original_name),"Cache-Control":"private, max-age=60"})


@router.delete("/equipment/{equipment_id}/documents/{document_id}", status_code=204, dependencies=[Depends(require_csrf)])
def delete_equipment_document(equipment_id:UUID,document_id:UUID,db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))]):
    row=db.get(Equipment,equipment_id)
    if not row: raise HTTPException(404,"Equipment not found")
    doc=db.get(EquipmentDocument,document_id)
    if not doc or doc.equipment_id!=row.id: raise HTTPException(404,"Document not found")
    path=doc.storage_path
    db.add(AuditLog(user_id=user.id,actor_name=user.display_name,action="Deleted equipment document",module="Equipment Register",entity_type="equipment",entity_id=str(row.id),metadata_json={"document_id":str(doc.id),"kind":doc.kind,"name":doc.original_name}))
    db.delete(doc); db.commit()
    if settings.storage_configured: SupabaseStorage().remove([path])
    return Response(status_code=204)


def _norm_header(value) -> str:
    return "_".join(str(value or "").strip().lower().replace("/"," ").replace("-"," ").split())


_IMPORT_ALIASES={
    "plant":{"plant","plant_code","site_code","site"},"equipment_code":{"equipment_code","code","tag","equipment_tag"},
    "equipment_name":{"equipment_name","name","equipment"},"equipment_type":{"equipment_type","type","category"},
    "maker":{"maker","manufacturer"},"model":{"model"},"serial_number":{"serial_number","serial","serial_no"},
    "status":{"status"},"discom":{"discom"},"feeder_or_gss":{"feeder_or_gss","feeder","gss"},"rating":{"rating"},
    "capacity_kw":{"capacity_kw","capacity"},"quantity":{"quantity","qty"},"location_in_plant":{"location_in_plant","location"},
    "manufacture_year":{"manufacture_year","year"},"commissioned_on":{"commissioned_on","commissioned"},
    "warranty_start":{"warranty_start"},"warranty_end":{"warranty_end"},"amc_vendor":{"amc_vendor"},"amc_end":{"amc_end"},
    "supplier":{"supplier"},"purchase_order_no":{"purchase_order_no","po_no","po"},"criticality":{"criticality"},"notes":{"notes"},
    "in_charge_user_id":{"in_charge_user_id"},"last_service_on":{"last_service_on"},"next_service_due":{"next_service_due"},
}


def _read_equipment_file(file: UploadFile) -> list[list]:
    raw=file.file.read()
    name=(file.filename or "").lower()
    if name.endswith(".csv") or (file.content_type or "").startswith("text/"):
        text=raw.decode("utf-8-sig",errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    try:
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); ws=wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    except Exception as exc:
        raise HTTPException(422,"Upload a valid .xlsx or .csv file") from exc


def _date_value(v):
    if v in (None,""): return None
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    try: return date.fromisoformat(str(v).strip()[:10])
    except Exception: return None


def _parse_import(db:Session,matrix:list[list]) -> tuple[list[dict],list[EquipmentCreate]]:
    if not matrix: return [],[]
    headers=[_norm_header(x) for x in matrix[0]]
    mapping={}
    for field,aliases in _IMPORT_ALIASES.items():
        for i,h in enumerate(headers):
            if h in aliases: mapping[field]=i; break
    plants=list(db.scalars(select(Plant)))
    plant_by={str(p.id).lower():p for p in plants}
    for p in plants:
        plant_by[p.site_code.strip().lower()]=p; plant_by[p.plant_name.strip().lower()]=p
    previews=[]; valid=[]
    for row_no,row in enumerate(matrix[1:],2):
        def cell(field):
            i=mapping.get(field); return row[i] if i is not None and i<len(row) else None
        if not any(v not in (None,"") for v in row): continue
        errors=[]; plant_key=str(cell("plant") or "").strip().lower(); plant=plant_by.get(plant_key)
        if not plant: errors.append("Plant/site code not found")
        name=str(cell("equipment_name") or "").strip()
        if not name: errors.append("Equipment name is required")
        status=str(cell("status") or "active").strip().lower()
        if status not in {"active","maintenance","fault","inactive"}: errors.append("Invalid status")
        criticality=str(cell("criticality") or "").strip().title() or None
        if criticality and criticality not in {"High","Medium","Low"}: errors.append("Invalid criticality")
        try: quantity=max(1,int(cell("quantity") or 1))
        except Exception: quantity=1; errors.append("Quantity must be a number")
        try: capacity=float(cell("capacity_kw")) if cell("capacity_kw") not in (None,"") else None
        except Exception: capacity=None; errors.append("Capacity must be a number")
        try: year=int(cell("manufacture_year")) if cell("manufacture_year") not in (None,"") else None
        except Exception: year=None; errors.append("Manufacture year must be a number")
        payload={
            "plant_id":plant.id if plant else uuid.uuid4(),"equipment_code":str(cell("equipment_code") or "").strip() or None,
            "equipment_name":name or "Invalid","equipment_type":str(cell("equipment_type") or "").strip() or None,
            "maker":str(cell("maker") or "").strip() or None,"model":str(cell("model") or "").strip() or None,
            "serial_number":str(cell("serial_number") or "").strip() or None,"status":status,"discom":str(cell("discom") or "").strip() or (plant.discom if plant else None),
            "feeder_or_gss":str(cell("feeder_or_gss") or "").strip() or None,"rating":str(cell("rating") or "").strip() or None,
            "capacity_kw":capacity,"quantity":quantity,"location_in_plant":str(cell("location_in_plant") or "").strip() or None,
            "manufacture_year":year,"commissioned_on":_date_value(cell("commissioned_on")),"warranty_start":_date_value(cell("warranty_start")),
            "warranty_end":_date_value(cell("warranty_end")),"amc_vendor":str(cell("amc_vendor") or "").strip() or None,"amc_end":_date_value(cell("amc_end")),
            "supplier":str(cell("supplier") or "").strip() or None,"purchase_order_no":str(cell("purchase_order_no") or "").strip() or None,
            "criticality":criticality,"notes":str(cell("notes") or "").strip() or None,"in_charge_user_id":cell("in_charge_user_id") or None,
            "last_service_on":_date_value(cell("last_service_on")),"next_service_due":_date_value(cell("next_service_due")),"metadata":{},
        }
        if payload["in_charge_user_id"]:
            try: payload["in_charge_user_id"]=UUID(str(payload["in_charge_user_id"]))
            except Exception: errors.append("Invalid in-charge user id"); payload["in_charge_user_id"]=None
        preview={"row":row_no,"valid":not errors,"errors":errors,"plant":plant.site_code if plant else str(cell("plant") or ""),"equipment_code":payload["equipment_code"],"equipment_name":name,"equipment_type":payload["equipment_type"],"maker":payload["maker"]}
        previews.append(preview)
        if not errors: valid.append(EquipmentCreate(**payload))
    return previews,valid


@router.get("/equipment/import/template.xlsx")
def equipment_import_template(db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(get_current_user)]):
    wb=Workbook(); ws=wb.active; ws.title="Equipment"
    ws.append(["plant","equipment_code","equipment_name","equipment_type","maker","model","serial_number","status","discom","feeder_or_gss","rating","capacity_kw","quantity","location_in_plant","manufacture_year","commissioned_on","warranty_start","warranty_end","amc_vendor","amc_end","supplier","purchase_order_no","criticality","notes","in_charge_user_id","last_service_on","next_service_due"])
    ws.append(["SITE-01","INV-01","Inverter 1","Inverter","Sungrow","SG110CX","SN-001","active","DISCOM","11kV Feeder","110 kW",110,1,"Block B","2025","2025-01-15","2025-01-15","2030-01-14","","","Supplier","PO-001","High","", "","2026-09-01","2026-10-01"])
    neutralise_formulas(wb); out=io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=equipment-register-template.xlsx"})


@router.post("/equipment/import", dependencies=[Depends(require_csrf)])
def import_equipment(
    db:Annotated[Session,Depends(get_db)],user:Annotated[User,Depends(require_permission("plants.manage"))],
    file:UploadFile=File(...),dry_run:bool=Form(True),
):
    previews,valid=_parse_import(db,_read_equipment_file(file))
    invalid=sum(1 for r in previews if not r["valid"])
    result={"filename":file.filename,"total":len(previews),"valid":len(valid),"invalid":invalid,"rows":previews}
    if dry_run:
        result["dry_run"]=True; return result
    if invalid:
        raise HTTPException(422,detail={"message":"Import has invalid rows. Run dry-run and correct them before commit.","rows":previews})
    created=EquipmentService(db).bulk_create(EquipmentBulkCreate(items=valid),user)
    return {**result,"dry_run":False,"created":_equipment_out_many(db,created)}


@router.get("/equipment/export.xlsx")
def export_equipment(
    db:Annotated[Session,Depends(get_db)], user:Annotated[User,Depends(get_current_user)], plant_id:UUID|None=None,
    q:str|None=None,equipment_type:str|None=None,maker:str|None=None,status:str|None=None,discom:str|None=None,
    in_charge_user_id:UUID|None=None,warranty:str|None=None,no_in_charge:bool=False,maker_missing:bool=False,serial_missing:bool=False,
):
    filters=_list_filters(plant_id,q,equipment_type,maker,status,discom,in_charge_user_id,warranty,no_in_charge,maker_missing,serial_missing)
    rows=EquipmentRepository(db).list(**filters,**_equipment_scope(db,user),offset=0,limit=10000)
    plant_ids={row.plant_id for row in rows}
    plants_by_id={
        plant.id:plant
        for plant in db.scalars(select(Plant).where(Plant.id.in_(plant_ids)))
    } if plant_ids else {}
    details=_equipment_out_many(db,rows)
    wb=Workbook(); ws=wb.active; ws.title="Equipment register"
    headers=["Plant","DISCOM","Code","Name","Type","Maker","Model","Serial","Status","In charge","Type owners","Rating","Capacity kW","Quantity","Location","Manufacture year","Commissioned","Warranty start","Warranty end","AMC vendor","AMC end","Supplier","PO No","Criticality","Last service","Next service","Notes"]
    ws.append(headers)
    for row,out in zip(rows,details):
        plant=plants_by_id.get(row.plant_id)
        ws.append([plant.site_code if plant else "",row.discom,row.equipment_code,row.equipment_name,row.equipment_type,row.maker,row.model,row.serial_number,row.status,out["in_charge_name"],", ".join(out["type_owners"]),row.rating,float(row.capacity_kw) if row.capacity_kw is not None else None,row.quantity,row.location_in_plant,row.manufacture_year,row.commissioned_on,row.warranty_start,row.warranty_end,row.amc_vendor,row.amc_end,row.supplier,row.purchase_order_no,row.criticality,row.last_service_on,row.next_service_due,row.notes])
    neutralise_formulas(wb); bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return StreamingResponse(bio,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=equipment-register.xlsx"})
