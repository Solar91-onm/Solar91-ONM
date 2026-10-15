from __future__ import annotations

import csv
import io
import re
from typing import Annotated
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.auth.emails import canonical_email, is_gmail_family
from app.core.config import get_settings
from app.core.database import get_db
from app.core.models import AuditLog, Equipment, Plant, User, UserEquipmentAssignment, UserEquipmentTypeAssignment, UserPlantAssignment
from app.core.security import require_csrf, require_permission
from app.plants.codes import normalize_site_code, plant_code_index
from app.users.router import PORTFOLIO_ROLES, _active_head_count, out

from app.notifications.mailer import send_invite_email

router = APIRouter(prefix="/api/users", tags=["Users"])
settings = get_settings()

# A person added by the O&M Head before signing in is stored as a pending user row whose
# external_subject is "invite:<email>". Clerk's first verified sign-in promotes the row.
INVITE_PREFIX = "invite:"
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

ROLE_ALIASES = {
    "o&m head": "O&M Head", "om head": "O&M Head", "head": "O&M Head", "admin": "O&M Head",
    "assignee": "Assignee", "o&m team": "Assignee", "om team": "Assignee", "team": "Assignee",
    "Administration": "Administration", "ceo": "Administration", "cto": "Administration", "ceo cto": "Administration", "ceo / cto": "Administration",
    "site engineer": "Site Engineer", "engineer": "Site Engineer", "se": "Site Engineer",
    "viewer": "Viewer",
}
HEADER_ALIASES = {
    "name": {"name", "full name", "person", "employee", "employee name"},
    "email": {"email", "email id", "e-mail", "mail", "email address"},
    "role": {"role", "designation", "access", "access level"},
    "plants": {"plants", "plant", "site code", "site codes", "sites", "plant code", "plant codes", "assigned plants"},
    "equipment": {"equipment", "equipment id", "equipment ids", "equipment code", "equipment codes",
                  "asset", "asset id", "asset code", "assets", "assigned equipment",
                  "equipment type", "equipment types", "type", "types"},
}


class InviteIn(BaseModel):
    name: str = Field(default="", max_length=255)
    email: str = Field(max_length=240)
    role: str
    plants: list[str] = Field(default_factory=list, max_length=500)  # site codes or plant names
    equipment: list[str] = Field(default_factory=list, max_length=2000)  # equipment TYPES, e.g. "Inverter"


def _norm_code(value) -> str:
    return normalize_site_code(str(value or "").split("(", 1)[0])


def _tokens(value) -> list[str]:
    return [x.strip() for x in re.split(r"[,;/\n|]+", str(value or "")) if x.strip()]


def _role(value) -> str | None:
    key = " ".join(str(value or "").strip().lower().replace("_", " ").split())
    return ROLE_ALIASES.get(key)


def _plant_index(db: Session):
    rows = list(db.scalars(select(Plant)))
    return plant_code_index(rows), {p.plant_name.strip().lower(): p for p in rows}


def _resolve_plants(tokens: list[str], index) -> list:
    by_code, by_name = index
    found, missing = [], []
    for t in tokens:
        if t.strip().lower() in {"all", "*"}:
            continue
        p = by_code.get(_norm_code(t)) or by_name.get(t.strip().lower())
        if p:
            if p.id not in [x.id for x in found]:
                found.append(p)
        else:
            missing.append(t)
    if missing:
        raise ValueError("Unknown plant: " + ", ".join(missing))
    return found


def _resolve_types(tokens: list[str]) -> list[str]:
    """Match tokens to the portal's equipment types (case-insensitive)."""
    from app.core.security import EQUIPMENT_TYPES
    by_key = {t.casefold(): t for t in EQUIPMENT_TYPES}
    found: list[str] = []
    missing: list[str] = []
    for token in tokens:
        match = by_key.get(token.strip().casefold())
        if match is None:
            missing.append(token.strip())
        elif match not in found:
            found.append(match)
    if missing:
        raise ValueError("Unknown equipment type: " + ", ".join(missing))
    return found


def _upsert_person(
    db: Session, actor: User, index, *, name: str, email: str, role: str,
    plant_tokens: list[str], equipment_tokens: list[str],
):
    """Create or update one person. Raises ValueError BEFORE changing anything if the row is invalid."""
    email = (email or "").strip()
    key = email.lower()
    if not EMAIL_RE.match(email) or len(email) > 240:
        raise ValueError("Invalid email address")
    if role not in {"O&M Head", "Assignee", "Administration", "Site Engineer", "Viewer"}:
        raise ValueError("Unknown role")
    if role == "Assignee" and plant_tokens:
        raise ValueError("Assignees are assigned equipment types, not plants")
    plants = _resolve_plants(plant_tokens, index) if role == "Site Engineer" else []
    if role in PORTFOLIO_ROLES or role == "Viewer":
        _resolve_plants(plant_tokens, index)  # still reject typos, but ignore the result
    if role == "Site Engineer" and len(plants) != 1:
        raise ValueError("A Site Engineer needs exactly one plant")
    if role != "Assignee" and equipment_tokens:
        raise ValueError("Equipment assignments are only valid for Assignees")
    equipment = _resolve_types(equipment_tokens) if role == "Assignee" else []

    existing = db.scalar(select(User).where(func.lower(User.email) == key))
    if existing is None and is_gmail_family(email):
        wanted = canonical_email(email)   # Gmail ignores dots and +tags: same mailbox = same person
        existing = next(
            (u for u in db.scalars(select(User).where(
                User.email.ilike("%@gmail.com") | User.email.ilike("%@googlemail.com")))
             if canonical_email(u.email) == wanted),
            None,
        )
    before = None
    if existing:
        if existing.id == actor.id and role != "O&M Head":
            raise ValueError("You cannot change your own O&M Head access")
        if existing.role == "O&M Head" and role != "O&M Head" and _active_head_count(db, exclude_id=existing.id) == 0:
            raise ValueError("At least one active O&M Head is required")
        before = {"role": existing.role,
                  "plant_ids": [str(x) for x in db.scalars(select(UserPlantAssignment.plant_id).where(UserPlantAssignment.user_id == existing.id))],
                  "equipment_types": sorted(db.scalars(select(UserEquipmentTypeAssignment.equipment_type).where(UserEquipmentTypeAssignment.user_id == existing.id)))}
        user = existing
        user.role = role
        if (user.external_subject or "").startswith(INVITE_PREFIX) and name:
            user.display_name = name.strip()
        status = "updated"
    else:
        user = User(
            auth_provider="clerk", external_subject=INVITE_PREFIX + canonical_email(email), email=email,
            display_name=(name or "").strip() or email, role=role, is_active=True,
            metadata_json={"invited_by": actor.display_name},
        )
        db.add(user)
        db.flush()
        status = "created"

    db.execute(delete(UserPlantAssignment).where(UserPlantAssignment.user_id == user.id))
    db.execute(delete(UserEquipmentAssignment).where(UserEquipmentAssignment.user_id == user.id))
    db.execute(delete(UserEquipmentTypeAssignment).where(UserEquipmentTypeAssignment.user_id == user.id))
    for p in plants:
        db.add(UserPlantAssignment(user_id=user.id, plant_id=p.id, assignment_role=role))
    for item in equipment:
        db.add(UserEquipmentTypeAssignment(user_id=user.id, equipment_type=item))
    db.add(AuditLog(
        user_id=actor.id, actor_name=actor.display_name,
        action="Pre-registered user" if status == "created" else "Changed user access",
        module="Users", entity_type="user", entity_id=str(user.id),
        before_state=before, after_state={"role": role, "plant_ids": [str(p.id) for p in plants],
                                          "equipment_types": list(equipment)},
        metadata_json={"target_email": email, "via": "invite/import"},
    ))
    return status, user


def _scope_text(db: Session, user: User) -> str:
    if user.role == "Assignee":
        types = sorted(db.scalars(select(UserEquipmentTypeAssignment.equipment_type).where(UserEquipmentTypeAssignment.user_id == user.id)))
        return (", ".join(types) + " \u2014 in every plant") if types else "Equipment to be assigned by your O&M Head"
    if user.role == "Site Engineer":
        names = list(db.scalars(select(Plant.plant_name).join(UserPlantAssignment, UserPlantAssignment.plant_id == Plant.id).where(UserPlantAssignment.user_id == user.id)))
        return names[0] if names else "Your plant"
    if user.role in {"O&M Head", "Administration"}:
        return "All plants"
    return "Access will be set by your O&M Head"


def _email_invite(db: Session, actor: User, user: User) -> tuple[bool, str]:
    if not user.email:
        return False, "No email address"
    return send_invite_email(to_email=user.email, name=user.display_name, role=user.role,
                             scope=_scope_text(db, user), invited_by=actor.display_name)


@router.post("/{user_id}/resend-invite", dependencies=[Depends(require_csrf)])
def resend_invite(
    user_id: UUID,
    db: Annotated[Session, Depends(get_db)],
    actor: Annotated[User, Depends(require_permission("users.manage"))],
):
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "User not found")
    if not (target.external_subject or "").startswith(INVITE_PREFIX):
        raise HTTPException(409, "This person has already signed in")
    sent, reason = _email_invite(db, actor, target)
    return {"email_sent": sent, "email_note": reason}


@router.post("/invite", status_code=201, dependencies=[Depends(require_csrf)])
def invite_person(
    data: InviteIn,
    db: Annotated[Session, Depends(get_db)],
    actor: Annotated[User, Depends(require_permission("users.manage"))],
):
    role = _role(data.role) or data.role
    try:
        status, user = _upsert_person(db, actor, _plant_index(db), name=data.name, email=data.email,
                                          role=role, plant_tokens=data.plants, equipment_tokens=data.equipment)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    db.commit()
    db.refresh(user)
    sent, reason = False, "Already signed in \u2014 no invite needed"
    if (user.external_subject or "").startswith(INVITE_PREFIX) and user.is_active:
        sent, reason = _email_invite(db, actor, user)
    return {**out(user), "status": status, "email_sent": sent, "email_note": reason}


def _read_matrix(filename: str, data: bytes) -> list[list]:
    if filename.lower().endswith(".csv"):
        try:
            return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        except Exception as exc:
            raise HTTPException(422, "CSV could not be read") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    except Exception as exc:
        raise HTTPException(422, "Workbook could not be read") from exc


@router.post("/import", dependencies=[Depends(require_csrf)])
async def import_people(
    db: Annotated[Session, Depends(get_db)],
    actor: Annotated[User, Depends(require_permission("users.manage"))],
    file: UploadFile = File(...),
):
    name = file.filename or ""
    if not name.lower().endswith((".xlsx", ".xlsm", ".csv")):
        raise HTTPException(415, "Upload an .xlsx, .xlsm or .csv file")
    data = await file.read()
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "File is too large")
    matrix = _read_matrix(name, data)
    if not matrix:
        raise HTTPException(422, "File is empty")
    headers = [" ".join(str(x or "").strip().lower().replace("_", " ").split()) for x in matrix[0]]
    col: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for i, h in enumerate(headers):
            if h in aliases:
                col[field] = i
                break
    missing = [f for f in ("email", "role") if f not in col]
    if missing:
        raise HTTPException(422, detail={"message": "Missing required column(s): " + ", ".join(missing),
                                         "found": headers})
    if len(matrix) - 1 > 500:
        raise HTTPException(422, "Too many rows (max 500 per import)")

    index = _plant_index(db)
    created = updated = 0
    errors: list[dict] = []
    seen: set[str] = set()
    new_people: list[User] = []

    def cell(row, field):
        i = col.get(field)
        return row[i] if i is not None and i < len(row) else None

    for n, row in enumerate(matrix[1:], start=2):
        email = str(cell(row, "email") or "").strip()
        if not email and not str(cell(row, "name") or "").strip():
            continue
        if email.lower() in seen:
            errors.append({"row": n, "email": email, "error": "Duplicate email in file"})
            continue
        seen.add(email.lower())
        role = _role(cell(row, "role"))
        if not role:
            errors.append({"row": n, "email": email, "error": f"Unknown role '{cell(row, 'role')}'"})
            continue
        try:
            status, _ = _upsert_person(db, actor, index, name=str(cell(row, "name") or ""), email=email,
                                       role=role, plant_tokens=_tokens(cell(row, "plants")),
                                       equipment_tokens=_tokens(cell(row, "equipment")))
        except ValueError as exc:
            errors.append({"row": n, "email": email, "error": str(exc)})
            continue
        created += status == "created"
        updated += status == "updated"
        if status == "created":
            new_people.append(_)
    db.commit()
    emailed = 0
    for person in new_people:
        db.refresh(person)
        emailed += bool(_email_invite(db, actor, person)[0])
    return {"filename": name, "created": created, "updated": updated, "errors": errors, "emailed": emailed}