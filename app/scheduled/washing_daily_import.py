"""Daily washing upload: record the tables cleaned per plant per day from one spreadsheet.

One row is one plant on one day: Site code, Date, Tables cleaned (0 when no washing was done) and optional
Remarks. A 0 is saved as a not-done day with the reason "Other"; what happened goes in Remarks. Rows follow the
same rules as the plant cycle sheet (``create_sheet_reports``): entries by the O&M Head or a washing Assignee
are saved as confirmed; a Site Engineer's entries (own plant, last 45 days) wait for review with one notification
per plant; a day that already has a pending or confirmed report is never added again; and every row is checked
before anything is saved (all or nothing).

The downloadable sheet for a month is pre-filled with one row per plant and day, so people only type numbers.
Days already saved show their values; uploading them again unchanged is harmless ("Already saved").
"""
from __future__ import annotations

import io
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Any
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.models import AuditLog, Plant, User, WashingConfig, WashingDailyReport
from app.core.security import user_scoped_plant_ids
from app.core.spreadsheet import cell_text, map_columns, neutralise_formulas, parse_sheet_date, read_upload_rows
from app.core.timeutil import business_today
from app.plants.codes import normalize_site_code, plant_code_index
from app.scheduled.washing_cycles import (
    MAX_BACKDATE_DAYS,
    MONTHLY_CYCLE_TARGET,
    _completing_row,
    _cycle_sequence,
    _day_label,
    _is_washing_manager,
    _month_key,
    _month_wash_summary,
    _notify_wash,
    _plant_table_count,
    _resequence_month,
)

MAX_DAILY_BYTES = 2 * 1024 * 1024
MAX_DAILY_ROWS = 3000
MAX_HISTORY_DAYS = 730  # the cycle sheet shows months up to two years back; older days could not be reviewed there
MAX_TABLES = 5000
SHEET = "Daily washing"
HEADERS = ("Site code", "Plant name", "Date", "Tables cleaned", "Remarks", "Saved in portal")
SUMMARY_HEADERS = ("Site code", "Plant name", "Tables per cycle", "Cycles completed (of 3)", "Cycle 1 done on",
                   "Cycle 2 done on", "Cycle 3 done on", "Tables cleaned (confirmed)", "Tables awaiting review",
                   "Times washed", "Extra washes after 3 cycles", "Extra tables after 3 cycles", "Days worked", "Days not done")
NOT_DONE_REASON = "Other"  # the sheet asks for no reason: a 0 day is saved with this one
ALIASES = {
    "site_code": {"Site code", "Plant code", "Site", "Code", "Site id"},
    "date": {"Date", "Report date", "Washing date", "Cleaning date", "Work date", "Day"},
    "tables": {"Tables cleaned", "Tables washed", "Tables", "Cleaned tables", "No of tables cleaned", "Number of tables cleaned"},
    "remarks": {"Remarks", "Remark", "Notes", "Note", "Comments", "Comment"},
}


def _require_reporter(db: Session, user: User) -> bool:
    """Returns whether the user is on the O&M team (entries saved as confirmed)."""
    manager = _is_washing_manager(db, user)
    if user.role != "Site Engineer" and not manager:
        raise HTTPException(403, "Only a Site Engineer or the O&M team can upload daily washing")
    return manager


def _allowed_plants(db: Session, user: User, manager: bool) -> set[UUID] | None:
    return None if manager else (user_scoped_plant_ids(db, user) or set())


def _describe(job_status: str, tables: int) -> str:
    return f"{tables} table{'s' if tables != 1 else ''}" if job_status == "Done" else "0 tables (not done)"


# ───────────────────────── template ─────────────────────────
def _template_month(month: str | None) -> tuple[date, date]:
    today = business_today()
    if not month:
        first = today.replace(day=1)
    else:
        try:
            first = date(int(month[:4]), int(month[5:7]), 1)
            if len(month) != 7 or month[4] != "-":
                raise ValueError
        except (TypeError, ValueError):
            raise HTTPException(422, "Month must look like 2026-10")
    last = first.replace(day=monthrange(first.year, first.month)[1])
    if first > today:
        raise HTTPException(422, "Choose the current month or an earlier one")
    if (today - last).days > MAX_HISTORY_DAYS:
        raise HTTPException(422, "Choose a month within the last two years")
    return first, min(last, today)


def daily_template_bytes(db: Session, user: User, month: str | None = None, plant_ids: list[str] | None = None) -> tuple[bytes, str]:
    """A month's sheet with one row per set-up plant and day (up to today). Saved days show their values."""
    manager = _require_reporter(db, user)
    first, end = _template_month(month)
    month_key = _month_key(first)
    stmt = select(Plant, WashingConfig).join(WashingConfig, WashingConfig.plant_id == Plant.id).order_by(Plant.site_code)
    found = [(plant, config) for plant, config in db.execute(stmt).all()]
    allowed = _allowed_plants(db, user, manager)
    if allowed is not None:
        found = [(plant, config) for plant, config in found if plant.id in allowed]
    if plant_ids is not None:
        requested: set[UUID] = set()
        for value in plant_ids:
            try:
                requested.add(UUID(value))
            except ValueError as exc:
                raise HTTPException(422, f"Invalid plant id: {value}") from exc
        found = [(plant, config) for plant, config in found if plant.id in requested]
    else:
        found = [(plant, config) for plant, config in found if config.active]
    found = [(plant, config) for plant, config in found if _plant_table_count(db, plant.id, config) > 0]
    days = (end - first).days + 1
    if len(found) * days > MAX_DAILY_ROWS:
        raise HTTPException(422, f"That would be more than {MAX_DAILY_ROWS} rows; choose fewer plants")

    saved: dict[tuple[UUID, date], WashingDailyReport] = {}
    month_reports: dict[UUID, list[WashingDailyReport]] = {}
    if found:
        for report in db.scalars(select(WashingDailyReport).where(
            WashingDailyReport.plant_id.in_([plant.id for plant, _config in found]),
            WashingDailyReport.cycle_month == month_key,
        ).order_by(WashingDailyReport.created_at)):
            month_reports.setdefault(report.plant_id, []).append(report)
            key = (report.plant_id, report.report_date)
            current = saved.get(key)
            # a pending or confirmed report wins over an older rejected one for the same day
            if current is None or report.review_status != "Rejected" or current.review_status == "Rejected":
                saved[key] = report

    workbook = Workbook()
    guide = workbook.active
    guide.title = "How to fill"
    for text in (
        "One row is one plant on one day. Site code and Date are already filled in for you.",
        "Type the number of tables cleaned that day in 'Tables cleaned'.",
        "If no washing was done, type 0 in 'Tables cleaned' and write what happened in 'Remarks'.",
        "Leave a row completely empty if you have nothing to report for that day; empty rows are skipped.",
        "Rows marked in 'Saved in portal' are already recorded. Change those days in the portal, not in this file.",
        f"Dates may be typed as DD-MM-YYYY. Site Engineers can enter the last {MAX_BACKDATE_DAYS} days for their own plant.",
        "Entries by the O&M Head or a washing Assignee are saved as confirmed; a Site Engineer's entries wait for review.",
        "You can also add your own rows (any plant and date) with the same columns.",
        "The 'Monthly summary' sheet is for reading only: cycles completed, tables cleaned and how many times the tables were washed.",
        "Washing done after all 3 cycles are complete (e.g. robot cleaning) is saved and counted under 'Extra washes after 3 cycles'.",
    ):
        guide.append([text])
    guide.column_dimensions["A"].width = 110

    sheet = workbook.create_sheet(SHEET)
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for plant, _config in found:
        for offset in range(days):
            day = date.fromordinal(first.toordinal() + offset)
            report = saved.get((plant.id, day))
            if report is not None and report.review_status != "Rejected":
                done = report.job_status == "Done"
                status = "Confirmed" if report.review_status == "Confirmed" else "Pending review"
                sheet.append([plant.site_code, plant.plant_name, day, report.tables_cleaned if done else 0,
                              report.remarks, status])
            else:
                sheet.append([plant.site_code, plant.plant_name, day, None, None,
                              "Rejected — enter again" if report is not None else None])
            sheet.cell(sheet.max_row, 3).number_format = "DD-MM-YYYY"
    sheet.auto_filter.ref = f"A1:F{max(1, sheet.max_row)}"
    for index, width in enumerate((14, 32, 14, 16, 48, 20), start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    last_row = max(sheet.max_row, 2) + 200
    tables = DataValidation(type="whole", operator="between", formula1="0", formula2=str(MAX_TABLES), allow_blank=True,
                            showErrorMessage=True, error="Type a whole number of tables")
    sheet.add_data_validation(tables)
    tables.add(f"D2:D{last_row}")

    summary = workbook.create_sheet("Monthly summary")
    summary.append(SUMMARY_HEADERS)
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.freeze_panes = "C2"
    for plant, config in found:
        info = _month_wash_summary(month_reports.get(plant.id, []), _plant_table_count(db, plant.id, config))
        done_on = [info["cycle_completed_on"].get(n) for n in range(1, MONTHLY_CYCLE_TARGET + 1)]
        summary.append([
            plant.site_code, plant.plant_name, info["tables_per_cycle"], info["cycles_completed"],
            *[date.fromisoformat(value) if value else None for value in done_on],
            info["confirmed_tables"], info["pending_tables"], info["times_washed"], info["extra_times"],
            info["extra_tables"], info["days_worked"], info["days_not_done"],
        ])
        for column in (5, 6, 7):
            summary.cell(summary.max_row, column).number_format = "DD-MM-YYYY"
    for index, width in enumerate((14, 30, 14, 14, 14, 14, 14, 18, 18, 14, 18, 18, 12, 14), start=1):
        summary.column_dimensions[summary.cell(1, index).column_letter].width = width
    workbook.active = 1
    neutralise_formulas(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), f"Daily_Washing_{month_key}.xlsx"


# ───────────────────────── upload ─────────────────────────
@dataclass
class _Planned:
    """A report (saved or about to be) used to work out cycle numbers before anything is written."""
    report_date: date
    order: tuple
    tables_cleaned: int
    cycle_no: int
    cycle_completed: bool = False
    confirmed: bool = True       # only confirmed reports complete a cycle; pending ones still fill it


def _whole_number(text: str) -> int:
    value = float(text.replace(",", ""))
    if not value.is_integer() or not 0 <= value <= MAX_TABLES:
        raise ValueError
    return int(value)


def import_daily(db: Session, user: User, filename: str, content: bytes, *, dry_run: bool) -> dict[str, Any]:
    manager = _require_reporter(db, user)
    headers, rows = read_upload_rows(filename, content, sheet=SHEET, max_bytes=MAX_DAILY_BYTES,
                                     max_rows=MAX_DAILY_ROWS, row_label="day rows")
    columns = map_columns(headers, ALIASES)
    if "site_code" not in columns or "date" not in columns:
        raise HTTPException(422, 'The file needs a "Site code" column and a "Date" column')
    if "tables" not in columns:
        raise HTTPException(422, 'The file needs a "Tables cleaned" column')

    allowed = _allowed_plants(db, user, manager)
    plants = plant_code_index(db.scalars(select(Plant)))
    today = business_today()
    records: list[dict[str, Any]] = []
    seen: dict[tuple[UUID, date], int] = {}
    blank = 0
    for number, values in rows:
        def get(field: str) -> Any:
            index = columns.get(field)
            return values[index] if index is not None and index < len(values) else None

        site_text, tables_text = cell_text(get("site_code")), cell_text(get("tables"))
        remarks = cell_text(get("remarks"))
        if not tables_text and not remarks:
            blank += 1  # a pre-filled day left empty
            continue
        errors: list[str] = []
        plant = None
        if not site_text:
            errors.append("Site code is required")
        else:
            plant = plants.get(normalize_site_code(site_text))
            if plant is None:
                errors.append(f"Site code {site_text} was not found")
            elif allowed is not None and plant.id not in allowed:
                errors.append(f"{plant.site_code} is outside your plant scope")
        day = None
        try:
            day = parse_sheet_date(get("date"))
            if day is None:
                errors.append("Date is required")
        except ValueError:
            errors.append(f"Date '{cell_text(get('date'))}' is not a date; use DD-MM-YYYY")
        if day is not None:
            if day > today:
                errors.append("The date cannot be in the future")
            elif not manager and (today - day).days > MAX_BACKDATE_DAYS:
                errors.append(f"A Site Engineer can enter the last {MAX_BACKDATE_DAYS} days only; ask the O&M Head for older dates")
            elif (today - day).days > MAX_HISTORY_DAYS:
                errors.append("Dates more than two years old cannot be uploaded")
        tables = 0
        if tables_text:
            try:
                tables = _whole_number(tables_text)
            except ValueError:
                errors.append(f"Tables cleaned must be a whole number from 0 to {MAX_TABLES}")
        elif not errors:
            errors.append("Enter the tables cleaned (0 if no washing was done)")
        if len(remarks) > 4000:
            errors.append("Remarks must be 4000 characters or fewer")
        if plant is not None and day is not None and not errors:
            key = (plant.id, day)
            if key in seen:
                errors.append(f"{plant.site_code} on {day.strftime('%d-%m-%Y')} appears more than once (also row {seen[key]})")
            else:
                seen[key] = number
        records.append({
            "row": number, "plant": plant, "site_code": plant.site_code if plant else site_text, "date": day,
            "job_status": "Done" if tables else "Not done", "tables": tables, "reason": None if tables else NOT_DONE_REASON,
            "remarks": remarks or None, "errors": errors, "action": "Problem" if errors else "New", "detail": "",
        })

    candidates = [r for r in records if not r["errors"]]
    plant_ids = sorted({r["plant"].id for r in candidates}, key=str)
    configs: dict[UUID, WashingConfig] = {}
    totals: dict[UUID, int] = {}
    existing: list[WashingDailyReport] = []
    if plant_ids:
        stmt = select(WashingConfig).where(WashingConfig.plant_id.in_(plant_ids)).order_by(WashingConfig.plant_id)
        if not dry_run:
            # Same lock as a cycle sheet save (plant order avoids deadlocks): concurrent entries for these plants wait.
            stmt = stmt.with_for_update()
        configs = {config.plant_id: config for config in db.scalars(stmt)}
        totals = {plant_id: _plant_table_count(db, plant_id, config) for plant_id, config in configs.items()}
        months = sorted({_month_key(r["date"]) for r in candidates})
        existing = list(db.scalars(select(WashingDailyReport).where(
            WashingDailyReport.plant_id.in_(plant_ids),
            WashingDailyReport.cycle_month.in_(months),
            WashingDailyReport.review_status.in_(["Pending", "Confirmed"]),
        )))
    taken = {(report.plant_id, report.report_date): report for report in existing}

    for record in candidates:
        plant, day = record["plant"], record["date"]
        if totals.get(plant.id, 0) < 1:
            record["errors"].append(f"Washing has not been set up for {plant.site_code}; set it up first")
            continue
        saved = taken.get((plant.id, day))
        if saved is None:
            continue
        # a saved not-done day matches a 0 whatever reason it was stored with
        same = saved.job_status == record["job_status"] and saved.tables_cleaned == record["tables"]
        status = "confirmed" if saved.review_status == "Confirmed" else "pending review"
        if same:
            record["action"] = "Unchanged"
            record["detail"] = f"Already saved ({status})" + ("; the remark change is ignored" if (saved.remarks or None) != record["remarks"] else "")
        else:
            record["errors"].append(
                f"Already saved for {_day_label(day)} as {_describe(saved.job_status, saved.tables_cleaned)} "
                f"({status}). Change that day in the plant's washing sheet in the portal")

    # Work out each day's cycle exactly as saving does (_cycle_sequence): every day of the month in date order fills
    # the cycles up to the plant's table count, so a back-filled day also moves the days after it.
    new_records = sorted((r for r in candidates if not r["errors"] and r["action"] == "New"),
                         key=lambda r: (str(r["plant"].id), r["date"]))
    state: dict[tuple[UUID, str], list[_Planned]] = {}
    for report in existing:
        state.setdefault((report.plant_id, report.cycle_month), []).append(_Planned(
            report.report_date, (0, report.created_at.timestamp() if report.created_at else 0.0),
            report.tables_cleaned, report.cycle_no, False, report.review_status == "Confirmed"))
    for sequence, record in enumerate(new_records):
        record["planned"] = _Planned(record["date"], (1, sequence), record["tables"], 1, confirmed=manager)
        state.setdefault((record["plant"].id, _month_key(record["date"])), []).append(record["planned"])
    extra_of: dict[int, bool] = {}
    for (plant_id, _month), rows_of_month in state.items():
        rows_of_month.sort(key=lambda row: (row.report_date, row.order))
        for row, (cycle_no, extra) in zip(rows_of_month, _cycle_sequence([(row.tables_cleaned, True) for row in rows_of_month], totals.get(plant_id, 0))):
            row.cycle_no = cycle_no
            extra_of[id(row)] = extra
        for cycle_no in range(1, MONTHLY_CYCLE_TARGET + 1):
            completing = _completing_row([row for row in rows_of_month if row.cycle_no == cycle_no and row.confirmed], totals.get(plant_id, 0))
            if completing is not None:
                completing.cycle_completed = True
    for record in new_records:
        record["cycle_no"] = record["planned"].cycle_no
        # all three cycles already complete (e.g. robot cleaning): the day is still saved, as extra work on cycle 3
        record["extra"] = extra_of[id(record["planned"])]
        if not manager:
            record["planned"] = None
    for record in new_records:
        if record["errors"]:
            continue
        record["detail"] = f"{_describe(record['job_status'], record['tables'])} · Cycle {record['cycle_no']}"
        if record.get("extra"):
            record["detail"] += " · extra washing (all 3 cycles already complete)"
        elif record.get("planned") is not None and record["planned"].cycle_completed:
            record["detail"] += " · completes the cycle"

    for record in records:
        if record["errors"]:
            record["action"] = "Problem"
    problems = [{"row": r["row"], "errors": r["errors"]} for r in records if r["errors"]]
    to_save = [r for r in new_records if not r["errors"]]
    result: dict[str, Any] = {
        "kind": "daily",
        "new": len(to_save),
        "unchanged": sum(1 for r in records if r["action"] == "Unchanged"),
        "blank": blank,
        "problems": problems,
        "plants": len({r["plant"].id for r in to_save}),
        "tables": sum(r["tables"] for r in to_save),
        "extra": sum(1 for r in to_save if r.get("extra")),
        "review_status": "Confirmed" if manager else "Pending",
        "saved": 0,
        "rows": [{
            "row": r["row"], "site_code": r["site_code"], "plant": r["plant"].plant_name if r["plant"] else "",
            "date": r["date"].isoformat() if r["date"] else "", "action": r["action"], "detail": r["detail"],
            "errors": r["errors"],
        } for r in records],
    }
    if dry_run or problems or not to_save:
        db.rollback()  # releases the plant locks; nothing was written
        return result

    try:
        now = datetime.now(timezone.utc)
        created: list[tuple[dict[str, Any], WashingDailyReport]] = []
        for record in to_save:
            report = WashingDailyReport(
                plant_id=record["plant"].id, cycle_month=_month_key(record["date"]), cycle_no=record["cycle_no"],
                report_date=record["date"], job_status=record["job_status"], tables_cleaned=record["tables"],
                cleaned_table_codes=[], remarks=record["remarks"], exception_reason=record["reason"],
                reported_by_id=user.id, reported_by=user.display_name,
            )
            if manager:
                report.review_status = "Confirmed"
                report.reviewed_by_id = user.id
                report.reviewed_by = user.display_name
                report.reviewed_at = now
                report.review_note = "Entered by O&M"
            db.add(report)
            created.append((record, report))
        db.flush()
        for plant_id, month in sorted({(str(r.plant_id), r.cycle_month) for _rec, r in created}):
            _resequence_month(db, UUID(plant_id), month, totals[UUID(plant_id)])
        for record, report in created:
            db.add(AuditLog(
                user_id=user.id, actor_name=user.display_name, action="Submitted daily washing report",
                module="Washing", entity_type="washing_daily_report", entity_id=str(report.id),
                metadata_json={"plant": record["plant"].site_code, "cycle_month": report.cycle_month, "cycle_no": report.cycle_no,
                               "report_date": report.report_date.isoformat(), "job_status": report.job_status,
                               "tables_cleaned": report.tables_cleaned, "exception_reason": report.exception_reason,
                               "cycle_completed": bool(report.cycle_completed), "source": "excel-upload", "row": record["row"]},
            ))
        if not manager:
            by_plant: dict[UUID, list[WashingDailyReport]] = {}
            for _record, report in created:
                by_plant.setdefault(report.plant_id, []).append(report)
            for plant_id, reports in by_plant.items():
                reports.sort(key=lambda report: report.report_date)
                plant = next(r["plant"] for r, _rep in created if r["plant"].id == plant_id)
                done = [report for report in reports if report.job_status == "Done"]
                not_done = len(reports) - len(done)
                span = _day_label(reports[0].report_date) if len(reports) == 1 else f"{_day_label(reports[0].report_date)} – {_day_label(reports[-1].report_date)}"
                _notify_wash(
                    db, plant, configs[plant_id],
                    event_type="washing_daily_report",
                    title=f"Washing sheet awaiting review — Cycle {', '.join(str(n) for n in sorted({r.cycle_no for r in reports}))}",
                    message=(f"{user.display_name} uploaded {len(reports)} day{'s' if len(reports) != 1 else ''} ({span}): "
                             f"{sum(report.tables_cleaned for report in done)} tables cleaned"
                             + (f", {not_done} day{'s' if not_done != 1 else ''} not done" if not_done else "")
                             + ". Review required."),
                    dedupe_key=f"washing-upload:{reports[0].id}",
                    extra_user_ids=[user.id],
                )
        db.add(AuditLog(
            user_id=user.id, actor_name=user.display_name, action="Uploaded daily washing", module="Washing",
            metadata_json={"saved": len(created), "unchanged": result["unchanged"], "blank": blank,
                           "plants": result["plants"], "review_status": result["review_status"]},
        ))
        db.commit()
    except Exception:
        db.rollback()
        raise
    result["saved"] = len(created)
    return result
