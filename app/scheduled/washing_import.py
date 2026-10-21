from __future__ import annotations

import csv
import io
from typing import Any
from uuid import UUID

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.spreadsheet import neutralise_formulas
from app.core.models import AuditLog, Plant, User, UserPlantAssignment, WashingConfig, WashingDailyReport, WashingHistory, WashingTable
from app.core.security import WASHING_TYPE, user_scoped_plant_ids
from app.plants.codes import find_plant_by_code, normalize_site_code, plant_code_index
from app.scheduled.washing_cycles import (
    _ensure_washing_assignee,
    _is_washing_manager,
    _set_table_count,
)

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 500
IMPORT_HEADERS = (
    "Site code",
    "Plant name",
    "Number of tables",
    "Washing assignee email",
    "Weekly target %",
    "Active",
    "Remarks",
)


def _require_manager(db: Session, user: User) -> None:
    if not _is_washing_manager(db, user):
        raise HTTPException(403, "Only the O&M team can manage washing setup imports")


def _scope_plants(db: Session, user: User, plants: list[Plant]) -> list[Plant]:
    if _is_washing_manager(db, user):
        return plants
    if user.role == "Site Engineer":
        allowed = user_scoped_plant_ids(db, user) or set()
        return [plant for plant in plants if plant.id in allowed]
    return []


def _configured_rows(db: Session, plants: list[Plant]) -> dict[UUID, tuple[WashingConfig | None, int, str | None]]:
    if not plants:
        return {}
    ids = [plant.id for plant in plants]
    configs = {row.plant_id: row for row in db.scalars(select(WashingConfig).where(WashingConfig.plant_id.in_(ids)))}
    counts = dict(db.execute(
        select(WashingTable.plant_id, func.count(WashingTable.id))
        .where(WashingTable.plant_id.in_(ids))
        .group_by(WashingTable.plant_id)
    ).all())
    for plant_id, row in configs.items():
        if row.tables_per_cycle:
            counts[plant_id] = row.tables_per_cycle
    assignees = {
        user.id: user.email for user in db.scalars(
            select(User).where(User.id.in_([c.assigned_user_id for c in configs.values() if c.assigned_user_id]))
        )
    }
    return {
        plant.id: (configs.get(plant.id), int(counts.get(plant.id, 0)), assignees.get(configs[plant.id].assigned_user_id) if configs.get(plant.id) and configs[plant.id].assigned_user_id else None)
        for plant in plants
    }


def _workbook_bytes(db: Session, plants: list[Plant], *, include_setup: bool) -> bytes:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "How to fill"
    for text in (
        "Use one row for each plant.",
        "Site code must match an existing plant in Solar91.",
        "For a new setup, enter the number of tables. Assignee and other settings can be added later.",
        "Weekly target is optional and must be from 1 to 100; the default is 85.",
        "When updating a setup, blank cells mean leave the saved value as it is.",
        "Remarks: the current situation at the plant (e.g. water tanker every 3 days, south gate locked on Sundays).",
    ):
        guide.append([text])
    guide.column_dimensions["A"].width = 95
    guide.freeze_panes = "A2"

    setup = workbook.create_sheet("Washing setup")
    setup.append(IMPORT_HEADERS)
    setup.freeze_panes = "A2"
    setup.auto_filter.ref = "A1:G1"
    configs = _configured_rows(db, plants)
    if include_setup:
        for plant in plants:
            config, count, email = configs[plant.id]
            setup.append([
                plant.site_code,
                plant.plant_name,
                count or None,
                email,
                config.compliance_threshold_percent if config else None,
                "Yes" if config is None or config.active else "No",
                config.notes if config else None,
            ])
    for index, width in enumerate((18, 32, 20, 34, 18, 14, 50), start=1):
        setup.column_dimensions[setup.cell(1, index).column_letter].width = width
    lists = workbook.create_sheet("Lists")
    lists.append(["Assignee emails", "Site codes"])
    assignees = list(db.scalars(
        select(User.email).where(
            User.role == "Assignee",
            User.is_active.is_(True),
            User.email.is_not(None),
        ).order_by(User.email)
    ))
    for row, email in enumerate(assignees[:MAX_IMPORT_ROWS], start=2):
        lists.cell(row, 1, email)
    for row, plant in enumerate(plants[:MAX_IMPORT_ROWS], start=2):
        lists.cell(row, 2, plant.site_code)
    lists.sheet_state = "hidden"
    table_count = DataValidation(type="whole", operator="between", formula1="1", formula2="5000", allow_blank=True)
    weekly_target = DataValidation(type="whole", operator="between", formula1="1", formula2="100", allow_blank=True)
    active = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    assignee = DataValidation(type="list", formula1=f"'Lists'!$A$2:$A${max(2, len(assignees) + 1)}", allow_blank=True)
    for validation in (table_count, weekly_target, active, assignee):
        setup.add_data_validation(validation)
    table_count.add(f"C2:C{MAX_IMPORT_ROWS + 1}")
    weekly_target.add(f"E2:E{MAX_IMPORT_ROWS + 1}")
    active.add(f"F2:F{MAX_IMPORT_ROWS + 1}")
    assignee.add(f"D2:D{MAX_IMPORT_ROWS + 1}")
    neutralise_formulas(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template_bytes(db: Session, user: User) -> bytes:
    _require_manager(db, user)
    plants = list(db.scalars(select(Plant).order_by(Plant.site_code).limit(MAX_IMPORT_ROWS)))
    return _workbook_bytes(db, plants, include_setup=False)


def export_bytes(db: Session, user: User, plant_ids: list[str] | None = None) -> bytes:
    plants = list(db.scalars(select(Plant).order_by(Plant.site_code)))
    if plant_ids is not None:
        requested: set[UUID] = set()
        for value in plant_ids:
            try:
                requested.add(UUID(value))
            except ValueError as exc:
                raise HTTPException(422, f"Invalid plant id: {value}") from exc
        found = {plant.id for plant in plants}
        if requested - found:
            raise HTTPException(404, "One or more plants were not found")
        plants = [plant for plant in plants if plant.id in requested]
    plants = _scope_plants(db, user, plants)
    if user.role != "Site Engineer" and not _is_washing_manager(db, user):
        raise HTTPException(403, "You cannot export washing setup")
    return _workbook_bytes(db, plants[:MAX_IMPORT_ROWS], include_setup=True)


def _read_rows(filename: str, content: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(413, "The file must be 2 MB or smaller")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "The CSV must use UTF-8 encoding") from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = [str(header or "").strip() for header in (reader.fieldnames or [])]
        rows = [(index + 2, dict(row)) for index, row in enumerate(reader)]
    elif suffix == "xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(422, "The Excel file could not be opened") from exc
        if "Washing setup" not in workbook.sheetnames:
            raise HTTPException(422, 'The workbook must contain a "Washing setup" sheet')
        sheet = workbook["Washing setup"]
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        headers = [str(value or "").strip() for value in (first or [])]
        rows = []
        for index, values in enumerate(iterator, start=2):
            rows.append((index, dict(zip(headers, values))))
        workbook.close()
    else:
        raise HTTPException(422, "Choose an .xlsx or .csv file")
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(422, f"The file can contain at most {MAX_IMPORT_ROWS} plant rows")
    normalized = [header.strip().casefold() for header in headers]
    if "site code" not in normalized:
        raise HTTPException(422, 'The file must have a "Site code" column')
    return headers, rows


def _cell(row: dict[str, Any], header: str) -> str:
    for key, value in row.items():
        if str(key or "").strip().casefold() == header.casefold():
            return str(value).strip() if value is not None else ""
    return ""


def _plant_for_site(db: Session, site_code: str) -> Plant | None:
    return find_plant_by_code(db, site_code)


def _reduction_problem(db: Session, plant: Plant, current: int, wanted: int) -> str | None:
    if wanted >= current:
        return None
    report_id = db.scalar(select(WashingDailyReport.id).where(WashingDailyReport.plant_id == plant.id).limit(1))
    if report_id:
        return "Table count cannot be reduced after washing reports have been recorded"
    table_ids = list(db.scalars(
        select(WashingTable.id).where(WashingTable.plant_id == plant.id)
        .order_by(WashingTable.table_code.desc()).limit(current - wanted)
    ))
    if table_ids and db.scalar(select(WashingHistory.id).where(WashingHistory.washing_table_id.in_(table_ids)).limit(1)):
        return "Table count cannot be reduced while removed tables have washing history"
    return None


def import_setup(db: Session, user: User, filename: str, content: bytes, *, dry_run: bool) -> dict[str, Any]:
    _require_manager(db, user)
    _headers, rows = _read_rows(filename, content)
    plants = plant_code_index(db.scalars(select(Plant)))
    configs = _configured_rows(db, list(plants.values()))
    used: set[str] = set()
    prepared: list[dict[str, Any]] = []
    problems: list[dict[str, Any]] = []
    for row_number, row in rows:
        if not any(str(value or "").strip() for value in row.values()):
            continue
        errors: list[str] = []
        site_code = _cell(row, "Site code")
        key = normalize_site_code(site_code)
        if not site_code:
            errors.append("Site code is required")
        elif key not in plants:
            errors.append(f"Site code {site_code} was not found")
        elif key in used:
            errors.append(f"Site code {site_code} appears more than once")
        else:
            used.add(key)
        plant = plants.get(key)
        config, current_count, current_email = configs.get(plant.id, (None, 0, None)) if plant else (None, 0, None)
        count_text = _cell(row, "Number of tables")
        email = _cell(row, "Washing assignee email")
        target_text = _cell(row, "Weekly target %")
        active_text = _cell(row, "Active")
        notes = _cell(row, "Remarks") or _cell(row, "Notes")  # older files call the column "Notes"
        is_new = config is None or current_count == 0
        table_count = current_count
        assignee = None
        threshold = config.compliance_threshold_percent if config else 85
        active = config.active if config else True
        next_notes = config.notes if config else None

        if count_text:
            try:
                numeric = float(count_text.replace(",", ""))
                if not numeric.is_integer() or not 1 <= numeric <= 5000:
                    raise ValueError
                table_count = int(numeric)
            except ValueError:
                errors.append("Number of tables must be a whole number from 1 to 5000")
        elif is_new:
            errors.append("Number of tables is required for a new setup")

        if email:
            assignee = db.scalar(select(User).where(
                func.lower(User.email) == email.casefold(),
                User.role == "Assignee",
                User.is_active.is_(True),
            ))
            if not assignee:
                errors.append(f"Assignee email {email} is not an active Assignee")

        if target_text:
            try:
                numeric = float(target_text)
                if not numeric.is_integer() or not 1 <= numeric <= 100:
                    raise ValueError
                threshold = int(numeric)
            except ValueError:
                errors.append("Weekly target must be a whole number from 1 to 100")
        if active_text:
            if active_text.casefold() not in {"yes", "no"}:
                errors.append("Active must be Yes or No")
            else:
                active = active_text.casefold() == "yes"
        if len(notes) > 500:
            errors.append("Remarks must be 500 characters or fewer")
        elif notes:
            next_notes = notes

        if plant and not errors:
            current_assignee_id = config.assigned_user_id if config else None
            same_assignee = (assignee.id if assignee else current_assignee_id) == current_assignee_id
            changed = (
                table_count != current_count
                or threshold != (config.compliance_threshold_percent if config else 85)
                or active != (config.active if config else True)
                or next_notes != (config.notes if config else None)
                or not same_assignee
            )
            action = "New" if is_new else "Changed" if changed else "Unchanged"
            prepared.append({
                "row": row_number, "plant": plant, "config": config, "table_count": table_count,
                "assignee": assignee, "threshold": threshold, "active": active, "notes": next_notes,
                "action": action,
            })
        else:
            problems.append({"row": row_number, "errors": errors})
            prepared.append({"row": row_number, "site_code": site_code, "action": "Problem", "errors": errors})

    counts = {
        "created": sum(1 for item in prepared if item["action"] == "New"),
        "updated": sum(1 for item in prepared if item["action"] == "Changed"),
        "unchanged": sum(1 for item in prepared if item["action"] == "Unchanged"),
        "skipped": 0,
        "problems": problems,
        "rows": [
            {
                "row": item["row"],
                "site_code": item.get("plant").site_code if item.get("plant") else item.get("site_code", ""),
                "plant": item.get("plant").plant_name if item.get("plant") else "",
                "action": item["action"],
                "errors": item.get("errors", []),
            }
            for item in prepared
        ],
    }
    if dry_run or problems:
        return counts

    try:
        for item in prepared:
            if item["action"] == "Unchanged":
                continue
            config = item["config"]
            if config is None:
                config = WashingConfig(plant_id=item["plant"].id, team=[], frequency_days=14)
            if item["assignee"]:
                _ensure_washing_assignee(db, user, item["assignee"])
                config.assigned_user_id = item["assignee"].id
            config.compliance_threshold_percent = item["threshold"]
            config.active = item["active"]
            config.notes = item["notes"]
            db.add(config)
            config.tables_per_cycle = item["table_count"]
        db.add(AuditLog(
            user_id=user.id,
            actor_name=user.display_name,
            action="Imported washing setup",
            module="Washing",
            metadata_json={"created": counts["created"], "updated": counts["updated"], "unchanged": counts["unchanged"]},
        ))
        db.commit()
    except Exception:
        db.rollback()
        raise
    return counts