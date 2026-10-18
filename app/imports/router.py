from __future__ import annotations
import csv
import io
import re
from datetime import date, datetime
from typing import Annotated
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.core.config import get_settings
from app.core.database import get_db
from app.core.models import Plant, User
from app.core.security import require_csrf, require_permission
from app.plants.codes import normalize_site_code, plant_code_index

router=APIRouter(prefix="/api/imports",tags=["Data / Imports"])
settings=get_settings()
ALIASES={
 "site_code":{"site code","site_code","sitecode","plant code","plant/site code","code"},
 "plant_name":{"plant name","plant_name","name","site name"},
 "project_company":{"project company","project_company","project company name","project","company","spv name","spv/project"},
 "district":{"district"},"discom":{"discom"},"location":{"location","city","village"},
 "dc_capacity_kw":{"dc capacity kw","dc_capacity_kw","dc capacity","capacitykw","capacity kw","dc capacity kwp","dc capacity (kwp)","dc kwp","dc (kwp)","installed dc capacity (kwp)"},
 "ac_capacity_kw":{"ac capacity kw","ac_capacity_kw","ac capacity","ac capacity (kw)","ac kw","ac (kw)","installed ac capacity (kw)"},
 "commissioned_on":{"commissioned on","commissioned_on","cod","commissioning date"},
 "latitude":{"latitude","lat"},"longitude":{"longitude","lng","lon"},
}

def norm(v): return " ".join(re.sub(r"[^a-z0-9]+", " ", str(v or "").strip().lower()).split())
def parse_float(v):
    if v in (None,""): return None
    try:return float(v)
    except Exception:return None
def parse_date(v):
    if not v:return None
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    try:return date.fromisoformat(str(v)[:10])
    except Exception:return None

@router.post("/plants",dependencies=[Depends(require_csrf)])
async def import_plants(file:UploadFile=File(...),db:Session=Depends(get_db),user:User=Depends(require_permission("plants.manage"))):
    if not file.filename or not file.filename.lower().endswith((".xlsx",".xlsm",".csv")): raise HTTPException(415,"Upload an .xlsx, .xlsm, or .csv plant information file")
    data=await file.read()
    if len(data)>settings.max_upload_mb*1024*1024: raise HTTPException(413,"Workbook is too large")
    if file.filename.lower().endswith(".csv"):
        try:
            matrix=list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        except Exception as exc:
            raise HTTPException(422,"CSV could not be read") from exc
        if not matrix: raise HTTPException(422,"Workbook is empty")
        headers=[norm(x) for x in matrix[0]]; rows=iter(matrix[1:])
    else:
        try: wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        except Exception as exc: raise HTTPException(422,"Workbook could not be read") from exc
        ws=wb.active; rows=ws.iter_rows(values_only=True)
        try: headers=[norm(x) for x in next(rows)]
        except StopIteration: raise HTTPException(422,"Workbook is empty")
    mapping={}
    for field,aliases in ALIASES.items():
        aliases={norm(alias) for alias in aliases}
        for i,h in enumerate(headers):
            if h in aliases: mapping[field]=i; break
    if "site_code" not in mapping: raise HTTPException(422,detail={"message":"Missing required column","required":["Site Code"],"found":headers})
    created=updated=0; errors=[]; seen=set()
    existing_by_code=plant_code_index(db.scalars(select(Plant)))
    for idx,row in enumerate(rows,start=2):
        def val(k): return row[mapping[k]] if k in mapping and mapping[k]<len(row) else None
        raw_site=str(val("site_code") or "").strip(); name=str(val("plant_name") or "").strip()
        if not raw_site and not name: continue
        if not raw_site: errors.append({"row":idx,"error":"Site Code is required"}); continue
        site=normalize_site_code(raw_site)
        if not site:
            errors.append({"row":idx,"site_code":raw_site,"error":"Site Code must contain letters or digits"}); continue
        key=site
        if key in seen: errors.append({"row":idx,"site_code":site,"error":"Duplicate Site Code in workbook"}); continue
        seen.add(key)
        lat=parse_float(val("latitude")); lon=parse_float(val("longitude"))
        if lat is not None and not -90<=lat<=90: errors.append({"row":idx,"site_code":site,"error":"Invalid latitude"}); continue
        if lon is not None and not -180<=lon<=180: errors.append({"row":idx,"site_code":site,"error":"Invalid longitude"}); continue
        p=existing_by_code.get(site); is_new=p is None
        if not p and not name: errors.append({"row":idx,"site_code":site,"error":"Plant Name is required for a new plant"}); continue
        if not p:
            p=Plant(site_code=site,plant_name=name,source="excel-import",created_by=user.id)
            existing_by_code[site]=p
        else:
            p.site_code=site
        if name: p.plant_name=name
        for field in ("project_company","district","discom","location"):
            value=val(field)
            if value is not None and str(value).strip(): setattr(p,field,str(value).strip())
        for field in ("dc_capacity_kw","ac_capacity_kw"):
            value=val(field)
            parsed=parse_float(value)
            if parsed is not None: setattr(p,field,parsed)
        commissioned=val("commissioned_on")
        parsed_commissioned=parse_date(commissioned)
        if parsed_commissioned is not None: p.commissioned_on=parsed_commissioned
        if lat is not None: p.latitude=lat
        if lon is not None: p.longitude=lon
        db.add(p); created+=int(is_new); updated+=int(not is_new)
    db.commit()
    if file.filename.lower().endswith((".xlsx",".xlsm")): wb.close()
    return {"filename":file.filename,"created":created,"updated":updated,"invalid_rows":errors,"mapped_columns":mapping}

LOCATION_ALIASES = {
    "label": {"plant name", "plant_name", "site", "site code", "site_code"},
    "city": {"city", "district"},
    "latitude": {"latitude", "lat"},
    "longitude": {"longitude", "lon", "lng", "long"},
}
def extract_site_code(value) -> str:
    text = str(value or "").strip()
    return text.split("(", 1)[0].strip()


@router.post("/locations", dependencies=[Depends(require_csrf)])
async def import_locations(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("plants.manage")),
):
    import csv

    filename = file.filename or ""
    lower = filename.lower()
    if not lower.endswith((".xlsx", ".xlsm", ".csv")):
        raise HTTPException(415, "Upload an .xlsx, .xlsm, or .csv location file")
    data = await file.read()
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "Location file is too large")

    if lower.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
            matrix = list(csv.reader(io.StringIO(text)))
        except Exception as exc:
            raise HTTPException(422, "CSV could not be read") from exc
    else:
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            matrix = [list(row) for row in wb.active.iter_rows(values_only=True)]
        except Exception as exc:
            raise HTTPException(422, "Workbook could not be read") from exc

    if not matrix:
        raise HTTPException(422, "Location file is empty")
    headers = [norm(x) for x in matrix[0]]
    mapping = {}
    for field, aliases in LOCATION_ALIASES.items():
        for i, header in enumerate(headers):
            if header in aliases:
                mapping[field] = i
                break
    missing = [label for key, label in (("label", "Plant Name / Site Code"), ("latitude", "Latitude"), ("longitude", "Longitude")) if key not in mapping]
    if missing:
        raise HTTPException(422, detail={"message": "Missing required location columns", "required": missing, "found": headers})

    plants = list(db.scalars(select(Plant)))
    by_norm = plant_code_index(plants)
    updated = 0
    unmatched: list[str] = []
    invalid: list[dict] = []
    duplicates: list[str] = []
    seen: set[str] = set()

    for row_no, row in enumerate(matrix[1:], start=2):
        def cell(key):
            i = mapping.get(key)
            return row[i] if i is not None and i < len(row) else None

        label = str(cell("label") or "").strip()
        if not label:
            continue
        code = normalize_site_code(extract_site_code(label))
        lat = parse_float(cell("latitude")); lon = parse_float(cell("longitude"))
        if cell("latitude") in (None, "") and cell("longitude") in (None, ""):
            continue
        if lat is None or not -90 <= lat <= 90 or lon is None or not -180 <= lon <= 180:
            invalid.append({"row": row_no, "label": label, "error": "Invalid latitude/longitude"})
            continue
        if code in seen:
            duplicates.append(label)
            continue
        seen.add(code)
        plant = by_norm.get(code)
        if not plant:
            unmatched.append(label)
            continue
        plant.latitude = lat
        plant.longitude = lon
        city = str(cell("city") or "").strip()
        if city:
            plant.location = city
        plant.source = "location-import"
        db.add(plant)
        updated += 1

    db.commit()
    return {
        "filename": filename,
        "total": max(0, len(matrix) - 1),
        "updated": updated,
        "unmatched": unmatched,
        "invalid_rows": invalid,
        "duplicates": duplicates,
        "mapped_columns": mapping,
    }
