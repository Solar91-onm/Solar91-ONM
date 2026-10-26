"""Site Issues upload: raise many tickets from one spreadsheet.

Each row becomes a ticket under the same rules as the Raise Issue form (``POST /api/issues``): plant scope for
Site Engineers, equipment-type scope for Assignees, an Assignee must be an active user with the Assignee role
(the O&M Head may extend their scope while assigning), tickets start as Open or Assigned and get an
"issue.created" history entry. Every row is checked first and nothing is saved while any row has a problem.
A row matching a ticket that is not yet closed (same plant, category, equipment and description) is skipped,
so uploading the same file twice does not raise duplicates. Closing or resolving still happens in the portal,
where evidence and root cause are recorded.
"""
from __future__ import annotations

import io
import logging
from datetime import date, datetime, time, timedelta, timezone
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.models import AuditLog, Equipment, Issue, IssueHistory, Plant, User, UserEquipmentTypeAssignment
from app.core.security import EQUIPMENT_TYPES, has_portfolio_scope, user_scoped_equipment_types, user_scoped_plant_ids
from app.core.spreadsheet import cell_text, map_columns, neutralise_formulas, parse_sheet_date, read_upload_rows
from app.issues.deadlines import deadline_hours
from app.issues.schemas import ISSUE_SEVERITIES, normalize_issue_severity
from app.issues.service import PLANT_TIMEZONE, IssueService, _lock_issue_ticket_numbers, issue_local_today, next_issue_ticket_no
from app.plants.codes import normalize_site_code, plant_code_index

logger = logging.getLogger(__name__)

MAX_ISSUE_BYTES = 2 * 1024 * 1024
MAX_ISSUE_ROWS = 500
SHEET = "Site issues"
HEADERS = ("Site code", "Plant name", "Issue category", "Problem", "Description", "Severity", "Equipment / unit",
           "Occurred on", "Target date", "Assignee email", "Maker", "Impact", "Remarks")
ALIASES = {
    "site_code": {"Site code", "Plant code", "Site", "Code"},
    "category": {"Issue category", "Category", "Vertical", "Equipment type", "Asset"},
    "problem": {"Problem", "Subcategory", "Sub category", "Fault"},
    "description": {"Description", "Issue details", "Details", "Issue description"},
    "severity": {"Severity", "Priority", "Severity / Priority"},
    "equipment": {"Equipment / unit", "Equipment", "Unit", "Equipment ref", "Equipment reference"},
    "occurred": {"Occurred on", "Occurred", "Date", "Issue date", "Reported on", "Start date"},
    "target": {"Target date", "Target", "ETC", "Deadline", "Due date"},
    "assignee": {"Assignee email", "Assignee", "Assigned to"},
    "maker": {"Maker", "Make", "Manufacturer"},
    "impact": {"Impact"},
    "remarks": {"Remarks", "Remark", "Current status", "Current situation", "Status remarks", "Notes", "Comments"},
}
_CATEGORIES = {value.casefold(): value for value in EQUIPMENT_TYPES}
_SEVERITY_ORDER = ("Critical", "Medium")


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def _scoped_plants(db: Session, user: User) -> set | None:
    return None if has_portfolio_scope(user) else (user_scoped_plant_ids(db, user) or set())


def _active_assignees(db: Session) -> list[User]:
    return list(db.scalars(select(User).where(User.role == "Assignee", User.is_active.is_(True)).order_by(User.display_name)))


def _deadline_text(hours: dict[str, int]) -> str:
    def span(h: int) -> str:
        return f"{h // 24} day{'s' if h // 24 != 1 else ''}" if h % 24 == 0 else f"{h} hours"
    return f"Critical {span(hours['Critical'])}, Medium {span(hours['Medium'])}"


def template_bytes(db: Session, user: User) -> bytes:
    allowed = _scoped_plants(db, user)
    plants = [plant for plant in db.scalars(select(Plant).order_by(Plant.site_code)) if allowed is None or plant.id in allowed]
    categories = sorted(user_scoped_equipment_types(db, user) or []) if user.role == "Assignee" else None
    category_list = [value for value in EQUIPMENT_TYPES if categories is None or value.casefold() in categories]
    assignees = [person for person in _active_assignees(db) if person.email]

    workbook = Workbook()
    guide = workbook.active
    guide.title = "How to fill"
    for text in (
        "One row is one issue (ticket). Fill the 'Site issues' sheet; the lists for drop-downs are on the 'Lists' sheet.",
        "Required: Site code, Issue category, Description and Severity (Critical or Medium; an old 'Low' is read as Medium).",
        "Occurred on: the day the fault happened (DD-MM-YYYY). Blank means today.",
        "Target date: when it should be fixed. Blank uses the usual deadline ("+_deadline_text(deadline_hours(db))+").",
        "Assignee email: optional. It must be an active Assignee; the ticket then starts as Assigned instead of Open.",
        "Equipment / unit: optional code or name of the unit, e.g. INV-03. A registered unit of the plant is linked automatically.",
        "Remarks: optional. What is going on right now (e.g. 'Vendor called, spare expected Friday'). It becomes the ticket's first message.",
        "Nothing is saved until every row is correct. A row that matches a ticket that is still open is skipped.",
        "Photos, progress, resolution and closing are added in the portal on each ticket.",
    ):
        guide.append([text])
    guide.column_dimensions["A"].width = 115

    sheet = workbook.create_sheet(SHEET)
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for index, width in enumerate((14, 30, 22, 24, 50, 12, 18, 14, 14, 30, 16, 30, 45), start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for column in ("H", "I"):
        for row in range(2, MAX_ISSUE_ROWS + 2):
            sheet[f"{column}{row}"].number_format = "DD-MM-YYYY"

    lists = workbook.create_sheet("Lists")
    lists.append(["Issue categories", "Severities", "Assignee emails", "Site codes", "Plant names"])
    for row, value in enumerate(category_list, start=2):
        lists.cell(row, 1, value)
    for row, value in enumerate(_SEVERITY_ORDER, start=2):
        lists.cell(row, 2, value)
    for row, person in enumerate(assignees, start=2):
        lists.cell(row, 3, person.email)
    for row, plant in enumerate(plants, start=2):
        lists.cell(row, 4, plant.site_code)
        lists.cell(row, 5, plant.plant_name)
    last = MAX_ISSUE_ROWS + 1
    validations = [
        (DataValidation(type="list", formula1=f"'Lists'!$A$2:$A${max(2, len(category_list) + 1)}", allow_blank=True), f"C2:C{last}"),
        (DataValidation(type="list", formula1='"Critical,Medium"', allow_blank=True), f"F2:F{last}"),
        (DataValidation(type="list", formula1=f"'Lists'!$C$2:$C${max(2, len(assignees) + 1)}", allow_blank=True), f"J2:J{last}"),
        (DataValidation(type="list", formula1=f"'Lists'!$D$2:$D${max(2, len(plants) + 1)}", allow_blank=True), f"A2:A{last}"),
    ]
    for validation, cells in validations:
        sheet.add_data_validation(validation)
        validation.add(cells)
    workbook.active = 1
    neutralise_formulas(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _occurred_at(day: date, now: datetime) -> datetime:
    """09:00 plant time on the day it happened (never later than now)."""
    stamp = datetime.combine(day, time(9, 0), tzinfo=PLANT_TIMEZONE).astimezone(timezone.utc)
    return min(stamp, now)


def _target_at(day: date) -> datetime:
    """End of the target day, plant time."""
    return datetime.combine(day, time(23, 59), tzinfo=PLANT_TIMEZONE).astimezone(timezone.utc)


def import_issues(db: Session, user: User, filename: str, content: bytes, *, dry_run: bool) -> dict[str, Any]:
    headers, rows = read_upload_rows(filename, content, sheet=SHEET, max_bytes=MAX_ISSUE_BYTES,
                                     max_rows=MAX_ISSUE_ROWS, row_label="issue rows")
    columns = map_columns(headers, ALIASES)
    missing = [label for field, label in (("site_code", "Site code"), ("category", "Issue category"),
                                          ("description", "Description"), ("severity", "Severity")) if field not in columns]
    if missing:
        raise HTTPException(422, "The file needs these columns: " + ", ".join(missing))
    if not dry_run:
        # Same lock as the Raise Issue form: ticket numbers and the duplicate check run one upload at a time.
        _lock_issue_ticket_numbers(db)

    plants = plant_code_index(db.scalars(select(Plant)))
    allowed = _scoped_plants(db, user)
    own_types = user_scoped_equipment_types(db, user) if user.role == "Assignee" else None
    assignees = _active_assignees(db)
    by_email = {person.email.casefold(): person for person in assignees if person.email}
    by_name: dict[str, list[User]] = {}
    for person in assignees:
        by_name.setdefault(_norm(person.display_name), []).append(person)
    owned: dict[Any, set[str]] = {}
    if assignees:
        for user_id, equipment_type in db.execute(select(UserEquipmentTypeAssignment.user_id, UserEquipmentTypeAssignment.equipment_type)
                                                  .where(UserEquipmentTypeAssignment.user_id.in_([p.id for p in assignees]))):
            owned.setdefault(user_id, set()).add(_norm(equipment_type))
    now = datetime.now(timezone.utc)
    today = issue_local_today()

    records: list[dict[str, Any]] = []
    for number, values in rows:
        def get(field: str) -> Any:
            index = columns.get(field)
            return values[index] if index is not None and index < len(values) else None

        errors: list[str] = []
        notes: list[str] = []
        site_text = cell_text(get("site_code"))
        plant = plants.get(normalize_site_code(site_text)) if site_text else None
        if not site_text:
            errors.append("Site code is required")
        elif plant is None:
            errors.append(f"Site code {site_text} was not found")
        elif allowed is not None and plant.id not in allowed:
            errors.append(f"{plant.site_code} is outside your plant scope")
        category_text = cell_text(get("category"))
        category = _CATEGORIES.get(category_text.casefold())
        if not category_text:
            errors.append("Issue category is required")
        elif category is None:
            errors.append(f"Issue category '{category_text}' is not a portal category (see the Lists sheet)")
        elif own_types is not None and category.casefold() not in own_types:
            errors.append(f"{category} is outside your assigned equipment types")
        description = " ".join(cell_text(get("description")).split())
        if not description:
            errors.append("Description is required")
        elif len(description) > 4000:
            errors.append("Description must be 4000 characters or fewer")
        severity_text = cell_text(get("severity"))
        severity = normalize_issue_severity(severity_text)
        if severity not in ISSUE_SEVERITIES:
            errors.append("Severity must be Critical or Medium")
        problem, equipment_ref = cell_text(get("problem")), cell_text(get("equipment"))
        maker, impact, remarks = cell_text(get("maker")), cell_text(get("impact")), cell_text(get("remarks"))
        for label, value, limit in (("Problem", problem, 160), ("Equipment / unit", equipment_ref, 255),
                                    ("Maker", maker, 160), ("Impact", impact, 4000), ("Remarks", remarks, 4000)):
            if len(value) > limit:
                errors.append(f"{label} must be {limit} characters or fewer")
        occurred = target = None
        try:
            occurred = parse_sheet_date(get("occurred"))
            if occurred is not None and occurred > today:
                errors.append("Occurred on cannot be in the future")
        except ValueError:
            errors.append(f"Occurred on '{cell_text(get('occurred'))}' is not a date; use DD-MM-YYYY")
        try:
            target = parse_sheet_date(get("target"))
            if target is not None and target < (occurred or today):
                errors.append("Target date cannot be before the day the issue occurred (today when Occurred on is blank)")
        except ValueError:
            errors.append(f"Target date '{cell_text(get('target'))}' is not a date; use DD-MM-YYYY")
        assignee = None
        assignee_text = cell_text(get("assignee"))
        if assignee_text:
            assignee = by_email.get(assignee_text.casefold())
            if assignee is None:
                matches = by_name.get(_norm(assignee_text), [])
                if len(matches) == 1:
                    assignee = matches[0]
                elif len(matches) > 1:
                    errors.append(f"More than one Assignee is called {assignee_text}; use their email")
                else:
                    errors.append(f"{assignee_text} is not an active Assignee")
            if assignee is not None and category is not None and category.casefold() not in owned.get(assignee.id, set()):
                if user.role == "O&M Head":
                    notes.append(f"{assignee.display_name} will be given the {category} scope")
                else:
                    errors.append(f"{assignee.display_name} is not assigned to the {category} category")
        records.append({
            "row": number, "plant": plant, "site_text": site_text, "category": category, "description": description,
            "severity": severity, "problem": problem or None, "equipment_ref": equipment_ref or None, "maker": maker or None,
            "impact": impact or None, "remarks": remarks or None, "occurred": occurred, "target": target, "assignee": assignee,
            "errors": errors, "notes": notes, "action": "New",
        })

    # Registered units of the plants in the file, to link "Equipment / unit" the way the Raise Issue form does.
    plant_ids = {r["plant"].id for r in records if r["plant"] is not None}
    units_by_code: dict[tuple, list[Equipment]] = {}
    units_by_name: dict[tuple, list[Equipment]] = {}
    open_tickets: dict[tuple, str] = {}
    if plant_ids:
        for unit in db.scalars(select(Equipment).where(Equipment.plant_id.in_(plant_ids), Equipment.status != "inactive")):
            if unit.equipment_code:
                units_by_code.setdefault((unit.plant_id, _norm(unit.equipment_code)), []).append(unit)
            units_by_name.setdefault((unit.plant_id, _norm(unit.equipment_name)), []).append(unit)
        for ticket_no, plant_id, vertical, equipment_ref, description in db.execute(
            select(Issue.ticket_no, Issue.plant_id, Issue.vertical, Issue.equipment_ref, Issue.description).where(
                Issue.plant_id.in_(plant_ids), Issue.deleted_at.is_(None), Issue.status != "Closed",
            )
        ):
            open_tickets.setdefault((plant_id, _norm(vertical), _norm(equipment_ref), _norm(description)), ticket_no)

    seen: dict[tuple, int] = {}
    for record in records:
        if record["errors"]:
            continue
        key = (record["plant"].id, _norm(record["category"]), _norm(record["equipment_ref"]), _norm(record["description"]))
        if key in seen:
            record["errors"].append(f"Same issue as row {seen[key]}")
            continue
        seen[key] = record["row"]
        if key in open_tickets:
            record["action"] = "Already open"
            record["notes"] = [f"Already raised as {open_tickets[key]}; skipped"]
            continue
        if record["equipment_ref"]:
            ref = _norm(record["equipment_ref"])
            units = units_by_code.get((record["plant"].id, ref)) or units_by_name.get((record["plant"].id, ref)) or []
            unit = units[0] if len(units) == 1 else None
            if unit is not None and (not unit.equipment_type or _norm(unit.equipment_type) == _norm(record["category"])):
                record["equipment_id"] = unit.id
                record["notes"].append(f"Linked to registered unit {unit.equipment_code or unit.equipment_name}")

    for record in records:
        if record["errors"]:
            record["action"] = "Problem"
    problems = [{"row": r["row"], "errors": r["errors"]} for r in records if r["errors"]]
    to_create = [r for r in records if r["action"] == "New"]
    result: dict[str, Any] = {
        "kind": "issues",
        "new": len(to_create),
        "unchanged": sum(1 for r in records if r["action"] == "Already open"),
        "problems": problems,
        "critical": sum(1 for r in to_create if r["severity"] == "Critical"),
        "assigned": sum(1 for r in to_create if r["assignee"] is not None),
        "saved": 0,
        "tickets": [],
        "rows": [{
            "row": r["row"], "site_code": r["plant"].site_code if r["plant"] else r["site_text"],
            "plant": r["plant"].plant_name if r["plant"] else "", "category": r["category"] or "",
            "description": r["description"][:120], "remarks": (r["remarks"] or "")[:120], "severity": r["severity"] or "",
            "assignee": r["assignee"].display_name if r["assignee"] else "", "action": r["action"],
            "detail": "; ".join(r["notes"]), "errors": r["errors"],
        } for r in records],
    }
    if dry_run or problems or not to_create:
        db.rollback()
        return result

    service = IssueService(db)
    # Same default deadlines as the Raise Issue form; the O&M Head sets them (app/issues/deadlines.py).
    hours = deadline_hours(db)
    try:
        number = int(next_issue_ticket_no(db)[4:])
        granted: set[tuple] = set()
        created: list[tuple[dict[str, Any], Issue]] = []
        for record in to_create:
            assignee = record["assignee"]
            if assignee is not None and record["category"].casefold() not in owned.get(assignee.id, set()) \
                    and (assignee.id, record["category"]) not in granted:
                granted.add((assignee.id, record["category"]))
                db.add(UserEquipmentTypeAssignment(user_id=assignee.id, equipment_type=record["category"]))
                db.add(AuditLog(user_id=user.id, actor_name=user.display_name, action="Equipment scope added while assigning a ticket",
                                module="Issues", entity_type="user", entity_id=str(assignee.id),
                                after_state={"equipment_type": record["category"], "assignee": assignee.display_name}))
            occurred_at = _occurred_at(record["occurred"], now) if record["occurred"] else now
            issue = Issue(
                ticket_no=f"ISS-{number:04d}", plant_id=record["plant"].id, equipment_id=record.get("equipment_id"),
                equipment_ref=record["equipment_ref"], vertical=record["category"], subcategory=record["problem"],
                description=record["description"], severity=record["severity"],
                status="Assigned" if assignee else "Open", occurred_at=occurred_at,
                etc_at=_target_at(record["target"]) if record["target"] else now + timedelta(hours=hours[record["severity"]]),
                assigned_to=assignee.id if assignee else None, assignee_text=assignee.display_name if assignee else None,
                maker=record["maker"], impact=record["impact"], photo_count=0, created_by=user.id,
                legacy_payload={"supplier": None, "serial_number": None, "evidence": []},
            )
            number += 1
            db.add(issue)
            created.append((record, issue))
        db.flush()
        for record, issue in created:
            service._history(issue, "issue.created", user, status_to=issue.status, metadata={"source": "excel-upload", "row": record["row"]})
            if record["remarks"]:
                # The current situation from the sheet opens the ticket's conversation, like a first update message.
                db.add(IssueHistory(issue_id=issue.id, action="issue.message", note=record["remarks"], actor_id=user.id,
                                    actor_name=user.display_name,
                                    metadata_json={"evidence_ids": [], "intent": "update", "role": user.role, "source": "excel-upload"}))
        db.add(AuditLog(user_id=user.id, actor_name=user.display_name, action="Uploaded site issues", module="Issues",
                        metadata_json={"created": len(created), "skipped_open": result["unchanged"],
                                       "first_ticket": created[0][1].ticket_no, "last_ticket": created[-1][1].ticket_no}))
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "Ticket numbers changed while saving; upload the file again") from exc
    except Exception:
        db.rollback()
        raise
    issues = [issue for _record, issue in created]
    _notify_upload(db, user, issues, {r["plant"].id: r["plant"].site_code for r, _issue in created})
    result["saved"] = len(issues)
    result["tickets"] = [issue.ticket_no for issue in issues]
    return result


def _notify_upload(db: Session, user: User, issues: list[Issue], site_codes: dict) -> None:
    """One summary per person instead of one alert per ticket; critical tickets still alert one by one.
    A failing notification never undoes the upload."""
    try:
        from app.notifications.service import equipment_responsible_users, notify_issue_assignee_and_heads, notify_users, users_with_role

        batch = str(issues[0].id)
        span = issues[0].ticket_no if len(issues) == 1 else f"{issues[0].ticket_no} – {issues[-1].ticket_no}"
        critical = [issue for issue in issues if issue.severity == "Critical"]
        notify_users(
            db, users_with_role(db, {"O&M Head"}), event_type="issue_created",
            title=f"{len(issues)} issue{'s' if len(issues) != 1 else ''} uploaded",
            message=f"{user.display_name} raised {span} from a spreadsheet"
                    + (f" · {len(critical)} critical" if critical else "")
                    + " · " + ", ".join(sorted(set(site_codes.values()))[:8]),
            target_type="issue", target_id=issues[0].ticket_no, dedupe_key=f"issue-upload:{batch}",
        )
        for issue in critical:
            notify_issue_assignee_and_heads(
                db, plant_id=issue.plant_id, assigned_to=issue.assigned_to, event_type="critical_issue",
                title="Critical issue raised — " + issue.ticket_no,
                message=f"{site_codes.get(issue.plant_id, '')} · {issue.description[:180]}",
                target_id=issue.ticket_no, include_heads=True, dedupe_key=f"issue-created:{issue.id}",
            )
        per_user: dict[Any, tuple[User, list[Issue]]] = {}
        owners_cache: dict[tuple, list[User]] = {}
        for issue in issues:
            if issue.severity == "Critical":
                continue
            if issue.assigned_to:
                people = [db.get(User, issue.assigned_to)]
            else:
                key = (issue.plant_id, issue.equipment_id, _norm(issue.equipment_ref), _norm(issue.vertical))
                if key not in owners_cache:
                    owners_cache[key] = equipment_responsible_users(db, plant_id=issue.plant_id, equipment_id=issue.equipment_id,
                                                                    equipment_ref=issue.equipment_ref, equipment_type=issue.vertical)
                people = owners_cache[key]
            for person in people:
                if person is not None and person.is_active and person.id != user.id:
                    per_user.setdefault(person.id, (person, []))[1].append(issue)
        for person, tickets in per_user.values():
            notify_users(
                db, [person], event_type="issue_created",
                title=f"{len(tickets)} new issue{'s' if len(tickets) != 1 else ''} for you",
                message=", ".join(t.ticket_no for t in tickets[:6]) + (f" and {len(tickets) - 6} more" if len(tickets) > 6 else "")
                        + f" · uploaded by {user.display_name}",
                target_type="issue", target_id=tickets[0].ticket_no, dedupe_key=f"issue-upload:{batch}:mine",
            )
        db.commit()
    except Exception:
        db.rollback()
        logger.exception("Issue upload notification failed; the tickets were created")
