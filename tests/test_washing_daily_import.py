from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.core.models import AuditLog, Notification, Plant, WashingDailyReport
from app.core.timeutil import business_today
from app.scheduled.schemas import WashingSheetEntry
from app.scheduled.washing_daily_import import NOT_DONE_REASONS
from tests import conftest
from tests.conftest import login_user
from tests.test_washing_cycles import BASE, _plant, _setup_cycle_users
from tests.test_washing_sheet import _days, _engineer

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _csv(text: str):
    return {"file": ("days.csv", text.encode("utf-8"), "text/csv")}


def _upload(client, files, dry_run):
    return client.post(f"{BASE}/daily-import?dry_run={'true' if dry_run else 'false'}", files=files)


def _d(day):
    return day.strftime("%d-%m-%Y")


def test_reasons_match_the_sheet_schema():
    assert "Rain" in NOT_DONE_REASONS
    for reason in NOT_DONE_REASONS:
        WashingSheetEntry(report_date=business_today(), job_status="Not done", exception_reason=reason)


def test_head_upload_previews_then_saves_confirmed_days_and_completes_cycles(client, db):
    plant = _plant(db)  # 4 tables per cycle
    head, _assignee = _setup_cycle_users(client, db, plant)
    conftest._CURRENT_USER = head
    d1, d2, d3, d4 = _days(4)
    text = ("Site code,Date,Tables cleaned,Not done reason,Remarks\n"
            f"cycle01,{_d(d1)},3,,\n"
            f"CYCLE01,{d2.isoformat()},,rain,Heavy rain\n"
            f"CYCLE01,{_d(d3)},1,,\n"
            f"CYCLE01,{_d(d4)},,,\n")  # a template day left empty

    preview = _upload(client, _csv(text), True)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert (body["new"], body["blank"], body["problems"], body["saved"]) == (3, 1, [], 0)
    assert body["review_status"] == "Confirmed" and body["tables"] == 4
    assert "completes the cycle" in body["rows"][2]["detail"]
    assert db.query(WashingDailyReport).count() == 0

    saved = _upload(client, _csv(text), False)
    assert saved.status_code == 200, saved.text
    assert saved.json()["saved"] == 3
    db.expire_all()
    reports = db.query(WashingDailyReport).order_by(WashingDailyReport.report_date).all()
    assert [r.review_status for r in reports] == ["Confirmed"] * 3
    assert [r.cycle_no for r in reports] == [1, 1, 1]
    assert reports[1].job_status == "Not done" and reports[1].exception_reason == "Rain" and reports[1].remarks == "Heavy rain"
    assert [r.cycle_completed for r in reports] == [False, False, True]
    assert db.query(AuditLog).filter_by(action="Uploaded daily washing").count() == 1

    # uploading the same file again changes nothing and is not a problem
    again = _upload(client, _csv(text), False).json()
    assert (again["new"], again["unchanged"], again["problems"], again["saved"]) == (0, 3, [], 0)
    assert db.query(WashingDailyReport).count() == 3

    # the next day's work counts towards cycle 2
    month_end_ok = d4.month == d3.month
    if month_end_ok:
        more = _upload(client, _csv(f"Site code,Date,Tables cleaned\nCYCLE01,{_d(d4)},2\n"), False).json()
        assert more["saved"] == 1
        db.expire_all()
        assert db.query(WashingDailyReport).filter_by(report_date=d4).one().cycle_no == 2


def test_any_problem_blocks_the_whole_file(client, db):
    plant = _plant(db)
    head, _assignee = _setup_cycle_users(client, db, plant)
    conftest._CURRENT_USER = head
    db.add(Plant(site_code="NOSETUP", plant_name="Not set up"))
    db.commit()
    d1, d2 = _days(2)
    future = business_today() + timedelta(days=1)
    text = ("Site code,Date,Tables cleaned,Not done reason\n"
            f"CYCLE01,{_d(d1)},2,\n"
            f"CYCLE01,{_d(d1)},1,\n"            # duplicate day
            f"MISSING,{_d(d2)},1,\n"            # unknown plant
            f"NOSETUP,{_d(d2)},1,\n"            # washing not set up
            f"CYCLE01,{_d(future)},1,\n"        # future
            f"CYCLE01,31-02-2026,1,\n"          # not a date
            f"CYCLE01,{_d(d2)},2,Rain\n"        # both
            f"CYCLE01,{_d(d2)},abc,\n")         # not a number
    response = _upload(client, _csv(text), False)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["saved"] == 0
    assert [p["row"] for p in body["problems"]] == [3, 4, 5, 6, 7, 8, 9]
    joined = " ".join(e for p in body["problems"] for e in p["errors"])
    for expected in ("more than once", "was not found", "has not been set up", "future", "not a date", "not both", "whole number"):
        assert expected in joined
    assert db.query(WashingDailyReport).count() == 0


def test_changed_value_for_a_saved_day_is_a_problem(client, db):
    plant = _plant(db)
    head, _assignee = _setup_cycle_users(client, db, plant)
    conftest._CURRENT_USER = head
    (d1,) = _days(1)
    assert _upload(client, _csv(f"Site code,Date,Tables cleaned\nCYCLE01,{_d(d1)},2\n"), False).json()["saved"] == 1
    body = _upload(client, _csv(f"Site code,Date,Tables cleaned\nCYCLE01,{_d(d1)},3\n"), False).json()
    assert body["saved"] == 0 and "Already saved" in body["problems"][0]["errors"][0]


def test_engineer_upload_is_pending_scoped_and_notifies_once(client, db):
    plant = _plant(db)
    head, assignee = _setup_cycle_users(client, db, plant)
    other = Plant(site_code="OTHER01", plant_name="Other plant")
    db.add(other)
    db.commit()
    engineer = _engineer(client, db, plant)
    d1, d2 = _days(2)
    old = business_today() - timedelta(days=60)

    blocked = _upload(client, _csv(f"Site code,Date,Tables cleaned\nOTHER01,{_d(d1)},1\nCYCLE01,{_d(old)},1\n"), False).json()
    assert "outside your plant scope" in blocked["problems"][0]["errors"][0]
    assert "last 45 days" in blocked["problems"][1]["errors"][0]

    before = db.query(Notification).filter_by(event_type="washing_daily_report").count()
    body = _upload(client, _csv(f"Site code,Date,Tables cleaned\nCYCLE01,{_d(d1)},1\nCYCLE01,{_d(d2)},2\n"), False).json()
    assert body["saved"] == 2 and body["review_status"] == "Pending"
    db.expire_all()
    assert {r.review_status for r in db.query(WashingDailyReport).all()} == {"Pending"}
    assert {r.reported_by_id for r in db.query(WashingDailyReport).all()} == {engineer.id}
    rows = db.query(Notification).filter_by(event_type="washing_daily_report").all()[before:]
    recipients = [n.user_id for n in rows]
    assert len(recipients) == len(set(recipients)) and head.id in recipients and assignee.id in recipients


def test_viewer_cannot_upload_or_download(client, db):
    login_user(client, db, role="Viewer", name="Viewer", email="daily-viewer@example.com")
    assert client.get(f"{BASE}/daily-template").status_code == 403
    assert _upload(client, _csv("Site code,Date,Tables cleaned\n"), True).status_code == 403


def test_month_template_is_prefilled_and_round_trips(client, db):
    plant = _plant(db)
    head, _assignee = _setup_cycle_users(client, db, plant)
    conftest._CURRENT_USER = head
    (d1,) = _days(1)
    month = d1.strftime("%Y-%m")
    assert _upload(client, _csv(f"Site code,Date,Tables cleaned\nCYCLE01,{_d(d1)},2\n"), False).json()["saved"] == 1

    response = client.get(f"{BASE}/daily-template?month={month}")
    assert response.status_code == 200, response.text
    assert f"Daily_Washing_{month}.xlsx" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    rows = list(book["Daily washing"].iter_rows(values_only=True))
    assert rows[0][:4] == ("Site code", "Plant name", "Date", "Tables cleaned")
    assert all(row[0] == "CYCLE01" for row in rows[1:])
    expected_days = (min(business_today(), (d1.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1))).day
    assert [row[2].day for row in rows[1:]] == list(range(1, expected_days + 1))
    saved_row = next(row for row in rows[1:] if row[2].date() == d1)
    assert saved_row[3] == 2 and saved_row[6] == "Confirmed"

    # filling one empty day in the downloaded sheet and uploading it saves only that day
    sheet = book["Daily washing"]
    target = next(index for index, row in enumerate(rows[1:], start=2) if row[2].date() != d1)
    sheet.cell(target, 4, 3)
    buffer = BytesIO()
    book.save(buffer)
    body = _upload(client, {"file": ("month.xlsx", buffer.getvalue(), XLSX)}, False).json()
    assert body["problems"] == [] and body["saved"] == 1 and body["unchanged"] == 1

    assert client.get(f"{BASE}/daily-template?month=2999-01").status_code == 422
    assert client.get(f"{BASE}/daily-template?month=bad").status_code == 422
