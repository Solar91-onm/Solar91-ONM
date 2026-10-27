from __future__ import annotations

from datetime import date, timedelta

from app.core.models import OMTask, Plant
from app.pmplan.cycle import CYCLE, CYCLE_DAYS, phase_for, phase_index, window
from tests.conftest import login_user

START = date(2026, 11, 1)


def _plant(db, code="KRN_01", name="Kerli"):
    p = Plant(site_code=code, plant_name=name)
    db.add(p); db.commit(); db.refresh(p)
    return p


def _body(**over):
    b = {"equipment_type": "Inverter", "task_name": "Inverter service", "frequency": "Monthly", "safety": "—", "est_hours": 2,
         "plants": ["KRN_01"], "start_date": START.isoformat(), "assignees": [], "per_unit": True, "follow_cycle": True}
    b.update(over)
    return b


def test_cycle_order_and_days_fill_thirty_days():
    days = {(p["area"], p["phase"]): (p["start_day"], p["end_day"]) for p in CYCLE}
    assert days[("Plant side", "Inverter")] == (1, 7)
    assert days[("Plant side", "ACCB")] == (8, 9)
    assert days[("Plant side", "VCB")] == (10, 11)
    assert days[("Plant side", "IDT Transformer")] == (12, 13)
    assert days[("Plant side", "Aux. Transformer + ACDB")] == (14, 15)
    assert days[("Plant side", "4 Pole")] == (16, 17)
    assert days[("Plant side", "Earth pit")] == (18, 19)
    assert days[("Plant side", "Miscellaneous")] == (20, 22)
    assert days[("Plant side", "Transmission line")] == (23, 24)
    assert days[("GSS side", "VCB")] == (25, 26)
    assert days[("GSS side", "4 Pole")] == (27, 28)
    assert days[("GSS side", "Transmission line")] == (29, 30)
    assert CYCLE[-1]["end_day"] == CYCLE_DAYS == 30
    assert phase_for("ACCB Panel")["equipment_types"] == ["ACCB Panel"]        # MCCB is no longer in PM
    assert phase_for("RMS")["phase"] == phase_for("CCTV")["phase"] == "Miscellaneous"
    assert phase_for("VCB Panel")["area"] == "Plant side" and phase_for("VCB - GSS side")["area"] == "GSS side"
    assert window("Earthing", START)["start"] == START + timedelta(days=17)


def test_inverter_units_spread_over_the_first_week(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    res = client.post("/api/pm-plan/tasks", json=_body(total_units=14))
    assert res.status_code == 200, res.text
    assert res.json()["cycle"]["start"] == START.isoformat() and res.json()["cycle"]["end"] == (START + timedelta(days=6)).isoformat()
    rows = db.query(OMTask).order_by(OMTask.unit_no).all()
    assert len(rows) == 14
    assert {t.due_on for t in rows} == {START + timedelta(days=i) for i in range(7)}       # 2 a day for 7 days
    assert all(t.start_on == START for t in rows)


def test_single_unit_equipment_uses_its_turn_in_the_cycle(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    for etype, start_day, end_day in (("VCB Panel", 10, 11), ("Transformer", 12, 13), ("VCB - GSS side", 25, 26),
                                      ("Transmission line - GSS side", 29, 30), ("RMS", 20, 22)):
        res = client.post("/api/pm-plan/tasks", json=_body(equipment_type=etype, task_name=f"{etype} PM"))
        assert res.status_code == 200, res.text
        task = db.query(OMTask).filter(OMTask.task_type == etype).one()
        assert task.start_on == START + timedelta(days=start_day - 1)
        assert task.due_on == START + timedelta(days=end_day - 1)      # one unit: due on the last day of its turn


def test_accb_1_and_2_are_one_job_and_template_apply_follows_cycle(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    res = client.post("/api/pm-plan/tasks", json=_body(equipment_type="ACCB Panel", task_name="ACCB PM", total_units=2,
                                                      units_per_day=2, save_template=True))
    assert res.status_code == 200, res.text
    task = db.query(OMTask).one()                                       # ACCB-1 + ACCB-2: one job, not two
    assert (task.unit_no, task.total_units, task.quantity) == (1, 1, 1)
    assert (task.start_on, task.due_on) == (START + timedelta(days=7), START + timedelta(days=8))
    _plant(db, "B_01", "Plant B")
    tpl = client.get("/api/pm-plan/templates").json()[0]
    res = client.post(f"/api/pm-plan/templates/{tpl['id']}/apply", json={"plants": ["B_01"], "start_date": START.isoformat(),
                                                                          "follow_cycle": True, "total_units": 2})
    assert res.status_code == 200 and res.json()["cycle"]["start_day"] == 8
    assert res.json()["count"] == 1


def test_miscellaneous_is_one_plant_side_job_after_earthing(client, db):
    assert phase_for("Miscellaneous")["phase"] == "Miscellaneous" and phase_for("Miscellaneous")["area"] == "Plant side"
    assert phase_index("Earthing") < phase_index("Miscellaneous") < phase_index("Transmission line")
    assert phase_for("Street Light") is phase_for("Miscellaneous")
    login_user(client, db, role="O&M Head")
    _plant(db)
    res = client.post("/api/pm-plan/tasks", json=_body(equipment_type="Miscellaneous", task_name="Street Light, ERT",
                                                      total_units=4))
    assert res.status_code == 200, res.text
    task = db.query(OMTask).one()
    assert (task.unit_no, task.total_units) == (1, 1)
    assert (task.start_on, task.due_on) == (START + timedelta(days=19), START + timedelta(days=21))     # cycle days 20-22


def test_miscellaneous_items_are_saved_as_the_remark(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    res = client.post("/api/pm-plan/tasks", json=_body(equipment_type="Miscellaneous", task_name="Miscellaneous PM",
                                                      remarks="Street Light, ERT"))
    assert res.status_code == 200, res.text
    assert "Remarks: Street Light, ERT" in [e["text"] for e in db.query(OMTask).one().log]


def test_aux_transformer_and_acdb_are_separate_single_jobs_sharing_two_days(client, db):
    assert phase_for("ACDB") is phase_for("Aux. Transformer") and phase_for("ACDB")["area"] == "Plant side"
    login_user(client, db, role="O&M Head")
    _plant(db)
    for etype in ("Aux. Transformer", "ACDB"):
        res = client.post("/api/pm-plan/tasks", json=_body(equipment_type=etype, task_name=f"{etype} PM", total_units=3))
        assert res.status_code == 200, res.text
        task = db.query(OMTask).filter(OMTask.task_type == etype).one()                 # one job each, not three
        assert (task.unit_no, task.total_units) == (1, 1)
        assert (task.start_on, task.due_on) == (START + timedelta(days=13), START + timedelta(days=14))   # cycle days 14-15
    assert phase_for("4 Pole")["start_day"] == 16                                       # later phases keep their days


def test_saved_scope_remembers_known_types(client, db):
    login_user(client, db, role="O&M Head")
    res = client.put("/api/pm-plan/scope", json={"sections": {"plant": ["Inverter"]}, "hidden": [],
                                                 "known": ["Inverter", "ACDB", "not a type"]})
    assert res.status_code == 200, res.text
    assert client.get("/api/pm-plan/scope").json()["known"] == ["Inverter", "ACDB"]


def test_mccb_is_removed_from_preventive_maintenance(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    res = client.post("/api/pm-plan/tasks", json=_body(equipment_type="MCCB", task_name="MCCB PM"))
    assert res.status_code == 422 and "no longer part of Preventive Maintenance" in res.text
    assert client.post("/api/pm-plan/templates", json={"equipment_type": "MCCB", "task_name": "MCCB PM"}).status_code == 422
    assert client.put("/api/pm-plan/unit-config", json={"plant": "KRN_01", "equipment_type": "MCCB", "total_units": 2,
                                                        "units_per_day": 1}).status_code == 422
    csv = "Equipment type,Plant(s),Task - what to do,Frequency\nMCCB,KRN_01,MCCB PM,Monthly\n".encode()
    up = client.post("/api/pm-plan/import?dry_run=true", files={"file": ("plan.csv", csv, "text/csv")})
    assert up.status_code == 422 and "no longer part of Preventive Maintenance" in up.text
    assert db.query(OMTask).count() == 0


def test_without_cycle_dates_are_unchanged(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    res = client.post("/api/pm-plan/tasks", json=_body(equipment_type="VCB Panel", follow_cycle=False))
    assert res.status_code == 200 and res.json()["cycle"] is None
    assert db.query(OMTask).one().due_on == START


def test_upload_sets_dates_from_the_cycle(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    csv = ("Equipment type,Plant(s),Task - what to do,Frequency,No. of units,Start date\n"
           "Inverter,KRN_01,Inverter service,Monthly,7,\n"
           "Earthing,KRN_01,Earth pit check,Monthly,,\n"
           "4 Pole - GSS side,KRN_01,GSS 4 pole check,Monthly,,\n"
           "VCB Panel,KRN_01,VCB service,Monthly,,01-12-2026\n").encode()
    q = f"follow_cycle=true&cycle_start={START.isoformat()}"
    pre = client.post(f"/api/pm-plan/import?dry_run=true&{q}", files={"file": ("plan.csv", csv, "text/csv")})
    assert pre.status_code == 200, pre.text
    sched = {s["equipment_type"]: s for s in pre.json()["schedule"]}
    assert sched["Earthing"]["start"] == "2026-11-18" and sched["Earthing"]["end"] == "2026-11-19"
    assert sched["4 Pole - GSS side"]["start"] == "2026-11-27"
    assert sched["VCB Panel"]["start"] == "2026-12-10"          # a row's own Start date is its cycle start
    assert db.query(OMTask).count() == 0
    done = client.post(f"/api/pm-plan/import?dry_run=false&{q}", files={"file": ("plan.csv", csv, "text/csv")})
    assert done.status_code == 200, done.text
    inv = db.query(OMTask).filter(OMTask.task_type == "Inverter").order_by(OMTask.unit_no).all()
    assert [t.due_on for t in inv] == [START + timedelta(days=i) for i in range(7)]
    assert db.query(OMTask).filter(OMTask.task_type == "Earthing").one().due_on == date(2026, 11, 19)


def test_cycle_endpoint_roles(client, db):
    login_user(client, db, role="O&M Head")
    body = client.get("/api/pm-plan/cycle").json()
    assert body["cycle_days"] == 30 and body["phases"][0]["equipment_types"] == ["Inverter"]
    login_user(client, db, role="Site Engineer", email="se@example.com")
    assert client.get("/api/pm-plan/cycle").status_code == 403


def test_chained_window_moves_later_phases_after_booked_work():
    from app.pmplan.cycle import chained_window, phase_index
    inv = phase_index("Inverter")
    assert chained_window("ACCB Panel", START)["start"] == START + timedelta(days=7)          # nothing booked: day 8
    w = chained_window("ACCB Panel", START, {inv: START + timedelta(days=9)})                 # inverters ran to day 10
    assert w["start"] == START + timedelta(days=10) and w["moved_days"] == 3
    gss = chained_window("Transmission line - GSS side", START, {inv: START + timedelta(days=9)})
    assert gss["start"] == START + timedelta(days=31) and gss["moved_days"] == 3              # the delay carries to the GSS side
    early = chained_window("ACCB Panel", START, {inv: START + timedelta(days=2)})             # finishing early keeps the standard day
    assert early["start"] == START + timedelta(days=7) and early["moved_days"] == 0


def test_accb_added_later_starts_after_the_plants_inverter_round(client, db):
    login_user(client, db, role="O&M Head")
    kerli = _plant(db)
    other = _plant(db, "B_01", "Plant B")
    # Kerli: 20 inverters at 2 a day = 10 days (Nov 1-10). Plant B has no inverter work booked.
    assert client.post("/api/pm-plan/tasks", json=_body(total_units=20, units_per_day=2)).status_code == 200
    res = client.post("/api/pm-plan/tasks", json=_body(equipment_type="ACCB Panel", task_name="ACCB PM", total_units=2,
                                                      plants=["KRN_01", "B_01"]))
    assert res.status_code == 200, res.text
    due = lambda plant: sorted(t.due_on for t in db.query(OMTask).filter(OMTask.task_type == "ACCB Panel",
                                                                          OMTask.plant_id == plant.id))
    assert due(kerli) == [date(2026, 11, 12)]          # one job, from the day after Kerli's last inverter (Nov 11-12)
    assert due(other) == [date(2026, 11, 9)]           # standard cycle days 8-9
    moved = res.json()["cycle_moved"]
    assert [m["plant"] for m in moved] == ["Kerli"] and moved[0]["moved_days"] == 3


def test_upload_rows_run_in_cycle_order_without_overlap(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    # rows out of order on purpose; the inverter round takes 10 days (20 units at 2 a day)
    csv = ("Equipment type,Plant(s),Task - what to do,Frequency,No. of units,Units per day\n"
           "VCB - GSS side,KRN_01,GSS VCB PM,Monthly,,\n"
           "VCB Panel,KRN_01,VCB service,Monthly,,\n"
           "ACCB Panel,KRN_01,ACCB PM,Monthly,2,\n"
           "Inverter,KRN_01,Inverter service,Monthly,20,2\n").encode()
    q = f"follow_cycle=true&cycle_start={START.isoformat()}"
    pre = client.post(f"/api/pm-plan/import?dry_run=true&{q}", files={"file": ("plan.csv", csv, "text/csv")})
    assert pre.status_code == 200, pre.text
    sched = {s["equipment_type"]: s for s in pre.json()["schedule"]}
    assert (sched["Inverter"]["start"], sched["Inverter"]["end"]) == ("2026-11-01", "2026-11-10")
    assert (sched["ACCB Panel"]["start"], sched["ACCB Panel"]["end"]) == ("2026-11-11", "2026-11-12")
    assert (sched["VCB Panel"]["start"], sched["VCB Panel"]["end"]) == ("2026-11-13", "2026-11-14")
    assert sched["VCB - GSS side"]["start"] == "2026-11-28" and sched["VCB - GSS side"]["moved_days"] == 3
    assert [s["equipment_type"] for s in pre.json()["schedule"]][:3] == ["Inverter", "ACCB Panel", "VCB Panel"]
    done = client.post(f"/api/pm-plan/import?dry_run=false&{q}", files={"file": ("plan.csv", csv, "text/csv")})
    assert done.status_code == 200, done.text
    vcb = db.query(OMTask).filter(OMTask.task_type == "VCB Panel").one()
    assert (vcb.start_on, vcb.due_on) == (date(2026, 11, 13), date(2026, 11, 14))


def _side_workbook(side: str, rows: list[tuple]) -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "PM Tasks"
    ws.append(["Equipment type", "Plant(s)", "Task - what to do", "Frequency - when"])
    for r in rows:
        ws.append(list(r))
    wb.properties.keywords = f"solar91-pm-side:{side}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_side_templates_reject_the_other_sides_equipment(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    q = f"follow_cycle=true&cycle_start={START.isoformat()}"
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    gss = _side_workbook("gss", [("VCB - GSS side", "KRN_01", "GSS VCB PM", "Monthly"), ("VCB Panel", "KRN_01", "VCB PM", "Monthly")])
    res = client.post(f"/api/pm-plan/import?dry_run=true&{q}", files={"file": ("gss.xlsx", gss, xlsx)})
    assert res.status_code == 422
    assert res.json()["detail"]["problems"] == [{"row": 3, "errors": [res.json()["detail"]["problems"][0]["errors"][0]]}]
    assert "not GSS side equipment" in res.json()["detail"]["problems"][0]["errors"][0]
    plant = _side_workbook("plant", [("4 Pole - GSS side", "KRN_01", "GSS 4 pole PM", "Monthly")])
    res = client.post(f"/api/pm-plan/import?dry_run=true&{q}", files={"file": ("plant.xlsx", plant, xlsx)})
    assert res.status_code == 422 and "GSS side template" in res.json()["detail"]["problems"][0]["errors"][0]
    ok = _side_workbook("gss", [("VCB - GSS side", "KRN_01", "GSS VCB PM", "Monthly")])
    assert client.post(f"/api/pm-plan/import?dry_run=true&{q}", files={"file": ("gss.xlsx", ok, xlsx)}).status_code == 200


def _filled_template(name: str, plant: str, inverters: int | None = None) -> bytes:
    import io
    from pathlib import Path
    from openpyxl import load_workbook
    wb = load_workbook(Path(__file__).resolve().parents[1] / "frontend" / "static" / "templates" / name)
    ws = wb["PM Tasks"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value:
            ws.cell(r, 2, plant)
            if inverters and ws.cell(r, 1).value == "Inverter":
                ws.cell(r, 9, inverters)
                ws.cell(r, 10, 2)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_plant_then_gss_templates_upload_in_order(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    q = f"follow_cycle=true&cycle_start={START.isoformat()}"
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    plant_file = _filled_template("PM_Plant_Side_Template.xlsx", "KRN_01")
    res = client.post(f"/api/pm-plan/import?dry_run=false&{q}", files={"file": ("plant.xlsx", plant_file, xlsx)})
    assert res.status_code == 200, res.text
    spans = [(s["equipment_type"], s["start"], s["end"]) for s in res.json()["schedule"]]
    assert spans[0] == ("Inverter", "2026-11-01", "2026-11-07")
    assert ("ACCB Panel", "2026-11-08", "2026-11-09") in spans and "MCCB" not in {s[0] for s in spans}
    assert spans[-2:] == [("Miscellaneous", "2026-11-20", "2026-11-22"), ("Transmission line", "2026-11-23", "2026-11-24")]
    assert db.query(OMTask).filter(OMTask.task_type == "ACCB Panel").count() == 1
    assert db.query(OMTask).filter(OMTask.task_type == "Miscellaneous").count() == 1
    assert ("ACDB", "2026-11-14", "2026-11-15") in spans and db.query(OMTask).filter(OMTask.task_type == "ACDB").count() == 1
    misc = db.query(OMTask).filter(OMTask.task_type == "Miscellaneous").one()
    assert "Remarks: Street Light, ERT" in [e["text"] for e in misc.log]
    assert all(a[2] < b[1] for a, b in zip(spans, spans[1:]) if a[1] != b[1])        # one equipment after the other
    gss = client.post(f"/api/pm-plan/import?dry_run=false&{q}",
                      files={"file": ("gss.xlsx", _filled_template("PM_GSS_Side_Template.xlsx", "KRN_01"), xlsx)})
    assert gss.status_code == 200, gss.text
    assert [(s["equipment_type"], s["start"], s["end"]) for s in gss.json()["schedule"]] == [
        ("VCB - GSS side", "2026-11-25", "2026-11-26"), ("4 Pole - GSS side", "2026-11-27", "2026-11-28"),
        ("Transmission line - GSS side", "2026-11-29", "2026-11-30")]


def test_long_inverter_round_moves_the_gss_side_upload_too(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db)
    q = f"follow_cycle=true&cycle_start={START.isoformat()}"
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # 20 inverters at 2 a day = 10 days, so everything after them moves 3 days
    res = client.post(f"/api/pm-plan/import?dry_run=false&{q}",
                      files={"file": ("plant.xlsx", _filled_template("PM_Plant_Side_Template.xlsx", "KRN_01", inverters=20), xlsx)})
    assert res.status_code == 200, res.text
    spans = {s["equipment_type"]: (s["start"], s["end"]) for s in res.json()["schedule"]}
    assert spans["Inverter"] == ("2026-11-01", "2026-11-10") and spans["ACCB Panel"] == ("2026-11-11", "2026-11-12")
    assert spans["Transmission line"] == ("2026-11-26", "2026-11-27")
    gss = client.post(f"/api/pm-plan/import?dry_run=true&{q}",
                      files={"file": ("gss.xlsx", _filled_template("PM_GSS_Side_Template.xlsx", "KRN_01"), xlsx)})
    first = gss.json()["schedule"][0]
    assert (first["equipment_type"], first["start"], first["moved_days"]) == ("VCB - GSS side", "2026-11-28", 3)
    assert gss.json()["schedule"][-1]["past_cycle_days"] == 3          # runs 3 days past the 30-day cycle
