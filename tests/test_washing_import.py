from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.core.models import AuditLog, Plant, UserEquipmentTypeAssignment, WashingConfig, WashingTable
from app.scheduled.washing_cycles import check_washing_overdue
from tests.conftest import login_user
from tests import conftest


BASE = "/api/scheduled/washing"


def _plant(db, site_code="IMP01", name="Import Plant"):
    plant = Plant(site_code=site_code, plant_name=name)
    db.add(plant)
    db.commit()
    db.refresh(plant)
    return plant


def _csv(rows: str):
    return {"file": ("setup.csv", rows.encode("utf-8"), "text/csv")}


def _head(client, db):
    return login_user(client, db, role="O&M Head", name="Import Head", email="import-head@example.com")


def test_template_and_export_include_current_setup_and_validation_lists(client, db):
    plant = _plant(db)
    assignee = login_user(client, db, role="Assignee", name="Plant Assignee", email="plant-assignee@example.com")
    _head(client, db)
    setup = client.post(f"{BASE}/setup", json={
        "plant_id": str(plant.id), "table_count": 8, "assignee_user_id": str(assignee.id),
        "compliance_threshold_percent": 92,
    })
    assert setup.status_code == 200, setup.text

    template = client.get(f"{BASE}/template")
    assert template.status_code == 200
    template_book = load_workbook(BytesIO(template.content), read_only=True)
    assert "Washing setup" in template_book.sheetnames and "How to fill" in template_book.sheetnames
    assert "plant-assignee@example.com" in [row[0] for row in template_book["Lists"].iter_rows(min_col=1, max_col=1, values_only=True)]

    exported = client.get(f"{BASE}/export")
    assert exported.status_code == 200
    export_book = load_workbook(BytesIO(exported.content), read_only=True)
    rows = list(export_book["Washing setup"].iter_rows(values_only=True))
    assert rows[1] == (plant.site_code, plant.plant_name, 8, assignee.email, 92, "Yes", None)
    reimport = client.post(f"{BASE}/import?dry_run=true", files={
        "file": ("roundtrip.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    })
    assert reimport.status_code == 200, reimport.text
    assert reimport.json()["unchanged"] == 1 and reimport.json()["problems"] == []


def test_import_dry_run_then_create_setup_and_write_audit(client, db):
    plant = _plant(db)
    assignee = login_user(client, db, role="Assignee", name="CSV Assignee", email="csv-assignee@example.com")
    head = _head(client, db)
    body = "Weekly target %,Site code,Notes,Washing assignee email,Active,Plant name,Number of tables\n90,IMP01,Use south gate,csv-assignee@example.com,Yes,ignored,11\n"

    preview = client.post(f"{BASE}/import?dry_run=true", files=_csv(body))
    assert preview.status_code == 200, preview.text
    assert preview.json()["created"] == 1 and preview.json()["problems"] == []
    assert db.get(WashingConfig, plant.id) is None

    saved = client.post(f"{BASE}/import?dry_run=false", files=_csv(body))
    assert saved.status_code == 200, saved.text
    assert saved.json()["created"] == 1
    config = db.get(WashingConfig, plant.id)
    assert config.assigned_user_id == assignee.id
    assert config.compliance_threshold_percent == 90
    assert config.active is True and config.notes == "Use south gate"
    assert db.query(WashingTable).filter_by(plant_id=plant.id).count() == 11
    assert db.query(UserEquipmentTypeAssignment).filter_by(user_id=assignee.id, equipment_type="Modules cleaning").count() == 1
    audit = db.query(AuditLog).filter_by(user_id=head.id, action="Imported washing setup").one()
    assert audit.metadata_json == {"created": 1, "updated": 0, "unchanged": 0}


def test_import_can_create_minimal_setup_without_assignee(client, db):
    plant = _plant(db)
    _head(client, db)
    response = client.post(
        f"{BASE}/import?dry_run=false",
        files=_csv("Site code,Number of tables\nIMP01,24\n"),
    )
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    config = db.get(WashingConfig, plant.id)
    assert config is not None and config.assigned_user_id is None
    assert config.compliance_threshold_percent == 85 and config.active is True
    assert db.query(WashingTable).filter_by(plant_id=plant.id).count() == 24


def test_import_updates_only_nonblank_fields_and_reports_unchanged(client, db):
    plant = _plant(db)
    assignee = login_user(client, db, role="Assignee", name="Existing Assignee", email="existing-assignee@example.com")
    head = _head(client, db)
    assert client.post(f"{BASE}/setup", json={
        "plant_id": str(plant.id), "table_count": 5, "assignee_user_id": str(assignee.id),
        "compliance_threshold_percent": 88,
    }).status_code == 200
    db.get(WashingConfig, plant.id).notes = "Saved notes"
    db.commit()

    unchanged = client.post(f"{BASE}/import?dry_run=true", files=_csv(
        "Site code,Plant name,Number of tables,Washing assignee email,Weekly target %,Active,Notes\n"
        "IMP01,ignored,5,,88,Yes,\n"
    ))
    assert unchanged.status_code == 200 and unchanged.json()["unchanged"] == 1

    changed = client.post(f"{BASE}/import?dry_run=false", files=_csv(
        "Site code,Plant name,Number of tables,Washing assignee email,Weekly target %,Active,Notes\n"
        "IMP01,ignored,7,,,No,\n"
    ))
    assert changed.status_code == 200 and changed.json()["updated"] == 1
    config = db.get(WashingConfig, plant.id)
    assert config.assigned_user_id == assignee.id and config.compliance_threshold_percent == 88
    assert config.active is False and config.notes == "Saved notes"
    assert db.query(WashingTable).filter_by(plant_id=plant.id).count() == 7


def test_import_problems_are_row_numbered_and_prevent_all_writes(client, db):
    plant = _plant(db)
    _head(client, db)
    body = (
        "Site code,Plant name,Number of tables,Washing assignee email,Weekly target %,Active,Notes\n"
        "IMP01,ignored,10,missing@example.com,88,Yes,\n"
        "MISSING,Unknown,5,missing@example.com,110,Maybe,\n"
    )
    result = client.post(f"{BASE}/import?dry_run=false", files=_csv(body))
    assert result.status_code == 200, result.text
    payload = result.json()
    assert payload["created"] == 0
    assert [problem["row"] for problem in payload["problems"]] == [2, 3]
    assert db.get(WashingConfig, plant.id) is None
    assert db.query(WashingTable).filter_by(plant_id=plant.id).count() == 0


def test_import_requires_washing_manager_and_enforces_file_limits(client, db):
    plant = _plant(db)
    engineer = login_user(client, db, role="Site Engineer", name="Plant Engineer", email="plant-engineer@example.com")
    from app.core.models import UserPlantAssignment
    db.add(UserPlantAssignment(user_id=engineer.id, plant_id=plant.id, assignment_role="Site Engineer"))
    db.commit()
    conftest._CURRENT_USER = engineer
    denied = client.post(f"{BASE}/import?dry_run=true", files=_csv(
        "Site code,Number of tables\nIMP01,5\n"
    ))
    assert denied.status_code == 403
    assert client.get(f"{BASE}/template").status_code == 403

    _head(client, db)
    too_large = client.post(
        f"{BASE}/import?dry_run=true",
        files={"file": ("setup.csv", b"x" * (2 * 1024 * 1024 + 1), "text/csv")},
    )
    assert too_large.status_code == 413


def test_inactive_setup_is_kept_but_does_not_raise_overdue_alert(client, db):
    plant = _plant(db)
    assignee = login_user(client, db, role="Assignee", name="Inactive Assignee", email="inactive-assignee@example.com")
    _head(client, db)
    assert client.post(f"{BASE}/setup", json={
        "plant_id": str(plant.id), "table_count": 4, "assignee_user_id": str(assignee.id),
    }).status_code == 200
    db.get(WashingConfig, plant.id).active = False
    db.commit()

    overview = client.get(f"{BASE}/overview").json()
    row = next(item for item in overview["plants"] if item["plant_id"] == str(plant.id))
    assert row["active"] is False and row["overdue_bucket"] == "on_track"
    assert overview["overdue_matrix"] == {"on_track": 0, "1d": 0, "3d": 0, "7d": 0}
    assert check_washing_overdue(db, date(2026, 10, 11)) == 0
