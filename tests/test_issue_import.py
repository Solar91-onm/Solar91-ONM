from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.core.models import Equipment, Issue, IssueHistory, Notification, Plant, UserEquipmentTypeAssignment, UserPlantAssignment
from app.core.timeutil import business_today
from tests import conftest
from tests.conftest import login_user

BASE = "/api/issues"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD = "Site code,Issue category,Problem,Description,Severity,Equipment / unit,Occurred on,Target date,Assignee email\n"


def _plant(db, code="ISSUP01", name="Upload Plant"):
    plant = Plant(site_code=code, plant_name=name)
    db.add(plant)
    db.commit()
    db.refresh(plant)
    return plant


def _csv(text):
    return {"file": ("issues.csv", text.encode("utf-8"), "text/csv")}


def _upload(client, text, dry_run):
    return client.post(f"{BASE}/import?dry_run={'true' if dry_run else 'false'}", files=_csv(text))


def test_head_uploads_tickets_with_assignee_scope_unit_link_and_no_duplicates(client, db):
    plant = _plant(db)
    unit = Equipment(plant_id=plant.id, equipment_code="INV-03", equipment_name="Inverter 3", equipment_type="Inverter")
    db.add(unit)
    db.commit()
    assignee = login_user(client, db, role="Assignee", name="Upload Assignee", email="upload-assignee@example.com")
    head = login_user(client, db, role="O&M Head", name="Upload Head", email="upload-head@example.com")
    occurred = (business_today() - timedelta(days=3)).strftime("%d-%m-%Y")
    text = HEAD + (
        f"issup01,inverter,Fault code,Inverter 3 tripping on earth fault,High,INV-03,{occurred},,upload-assignee@example.com\n"
        "ISSUP01,Vegetation,,Grass growth near row 12,Low,,,,\n"
    )

    preview = _upload(client, text, True)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert (body["new"], body["problems"], body["critical"], body["assigned"]) == (2, [], 1, 1)
    assert "given the Inverter scope" in body["rows"][0]["detail"] and "INV-03" in body["rows"][0]["detail"]
    assert db.query(Issue).count() == 0

    saved = _upload(client, text, False)
    assert saved.status_code == 200, saved.text
    assert saved.json()["tickets"] == ["ISS-0001", "ISS-0002"]
    db.expire_all()
    first, second = db.query(Issue).order_by(Issue.ticket_no).all()
    assert (first.severity, first.status, first.assigned_to, first.equipment_id, first.vertical, first.subcategory) == (
        "Critical", "Assigned", assignee.id, unit.id, "Inverter", "Fault code")
    assert (second.status, second.assigned_to, second.etc_at is not None) == ("Open", None, True)
    assert db.query(IssueHistory).filter_by(action="issue.created").count() == 2
    assert db.query(UserEquipmentTypeAssignment).filter_by(user_id=assignee.id, equipment_type="Inverter").count() == 1
    assert db.query(Notification).filter_by(user_id=assignee.id, event_type="critical_issue").count() == 1
    assert db.query(Notification).filter_by(user_id=head.id, title="2 issues uploaded").count() == 1

    again = _upload(client, text, False).json()
    assert (again["new"], again["unchanged"], again["saved"]) == (0, 2, 0)
    assert "ISS-0001" in again["rows"][0]["detail"]
    assert db.query(Issue).count() == 2


def test_problems_block_every_row(client, db):
    _plant(db)
    login_user(client, db, role="O&M Head", name="Upload Head", email="upload-head2@example.com")
    future = (business_today() + timedelta(days=2)).strftime("%d-%m-%Y")
    text = HEAD + (
        "ISSUP01,Inverter,,Valid row,Medium,,,,\n"
        "NOPE,Inverter,,Unknown plant,Medium,,,,\n"
        "ISSUP01,Spaceship,,Bad category,Medium,,,,\n"
        "ISSUP01,Inverter,,,Medium,,,,\n"
        "ISSUP01,Inverter,,Bad severity,Urgent,,,,\n"
        f"ISSUP01,Inverter,,Future,Low,,{future},,\n"
        "ISSUP01,Inverter,,Unknown assignee,Low,,,,nobody@example.com\n"
        "ISSUP01,Inverter,,Valid row,Medium,,,,\n"
    )
    body = _upload(client, text, False).json()
    assert body["saved"] == 0
    assert [p["row"] for p in body["problems"]] == [3, 4, 5, 6, 7, 8, 9]
    joined = " ".join(e for p in body["problems"] for e in p["errors"])
    for expected in ("not found", "not a portal category", "Description is required", "Critical or Medium",
                     "future", "not an active Assignee", "Same issue as row 2"):
        assert expected in joined
    assert db.query(Issue).count() == 0


def test_scope_rules_for_engineer_and_assignee(client, db):
    plant = _plant(db)
    other = _plant(db, "ISSUP02", "Other Plant")
    engineer = login_user(client, db, role="Site Engineer", name="Upload Engineer", email="upload-se@example.com")
    db.add(UserPlantAssignment(user_id=engineer.id, plant_id=plant.id, assignment_role="Site Engineer"))
    db.commit()
    body = _upload(client, HEAD + "ISSUP02,Inverter,,Other plant fault,Low,,,,\nISSUP01,Inverter,,My plant fault,Low,,,,\n", False).json()
    assert "outside your plant scope" in body["problems"][0]["errors"][0] and body["saved"] == 0
    assert _upload(client, HEAD + "ISSUP01,Inverter,,My plant fault,Low,,,,\n", False).json()["saved"] == 1

    worker = login_user(client, db, role="Assignee", name="Upload Worker", email="upload-worker@example.com")
    db.add(UserEquipmentTypeAssignment(user_id=worker.id, equipment_type="Transformer"))
    db.commit()
    conftest._CURRENT_USER = worker
    body = _upload(client, HEAD + f"{other.site_code},Inverter,,Not my type,Low,,,,\n", True).json()
    assert "outside your assigned equipment types" in body["problems"][0]["errors"][0]
    assert _upload(client, HEAD + f"{other.site_code},Transformer,,Oil leak,Medium,,,,\n", False).json()["saved"] == 1

    login_user(client, db, role="Administration", name="Upload Admin", email="upload-admin@example.com")
    assert client.get(f"{BASE}/import/template").status_code == 403
    assert _upload(client, HEAD, True).status_code == 403


def test_template_lists_categories_and_round_trips(client, db):
    _plant(db)
    login_user(client, db, role="O&M Head", name="Upload Head", email="upload-head3@example.com")
    response = client.get(f"{BASE}/import/template")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book["Site issues"]["A1"].value == "Site code"
    assert "Inverter" in [row[0] for row in book["Lists"].iter_rows(min_row=2, max_col=1, values_only=True)]
    sheet = book["Site issues"]
    for column, value in enumerate(("ISSUP01", None, "Module", "Hotspot", "Hotspot on string 4 modules", "Medium"), start=1):
        sheet.cell(2, column, value)
    buffer = BytesIO()
    book.save(buffer)
    saved = client.post(f"{BASE}/import?dry_run=false", files={"file": ("issues.xlsx", buffer.getvalue(), XLSX)})
    assert saved.status_code == 200, saved.text
    assert saved.json()["saved"] == 1
