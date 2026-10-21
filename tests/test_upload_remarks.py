"""Every upload sheet carries a Remarks column for the current situation, and the remark lands on the saved record."""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl import load_workbook

from app.core.models import Issue, IssueHistory, OMTask, Plant, WashingConfig
from tests.conftest import login_user


def _plant(db, code, name):
    plant = Plant(site_code=code, plant_name=name)
    db.add(plant)
    db.commit()
    db.refresh(plant)
    return plant


def test_issue_upload_remarks_become_the_first_ticket_message(client, db):
    _plant(db, "REM01", "Remarks Plant")
    login_user(client, db, role="O&M Head", name="Remarks Head", email="remarks-head@example.com")
    template = load_workbook(BytesIO(client.get("/api/issues/import/template").content))
    assert template["Site issues"].cell(1, 13).value == "Remarks"

    text = ("Site code,Issue category,Description,Severity,Remarks\n"
            "REM01,Inverter,INV-02 tripping,Medium,Vendor called; spare expected Friday\n"
            "REM01,Vegetation,Grass near row 4,Low,\n")
    preview = client.post("/api/issues/import?dry_run=true", files={"file": ("i.csv", text.encode(), "text/csv")}).json()
    assert preview["rows"][0]["remarks"] == "Vendor called; spare expected Friday"
    saved = client.post("/api/issues/import?dry_run=false", files={"file": ("i.csv", text.encode(), "text/csv")})
    assert saved.status_code == 200, saved.text
    first = db.query(Issue).filter_by(description="INV-02 tripping").one()
    messages = db.query(IssueHistory).filter_by(issue_id=first.id, action="issue.message").all()
    assert [m.note for m in messages] == ["Vendor called; spare expected Friday"]
    second = db.query(Issue).filter_by(description="Grass near row 4").one()
    assert db.query(IssueHistory).filter_by(issue_id=second.id, action="issue.message").count() == 0


def test_pm_plan_remarks_are_logged_on_every_created_task(client, db):
    login_user(client, db, role="O&M Head")
    _plant(db, "PMR-A", "PM Remarks Plant")
    payload = ("Equipment type,Plant(s),Task - what to do,Frequency,No. of units,Units per day,Start date,Remarks\r\n"
               'Inverter,PMR-A,Clean fans,Monthly,2,1,2026-11-02,"Filters ordered, INV-04 fan noisy"\r\n').encode()
    done = client.post("/api/pm-plan/import?dry_run=false", files={"file": ("plan.csv", payload, "text/csv")})
    assert done.status_code == 200, done.text
    tasks = db.query(OMTask).all()
    assert len(tasks) == 2
    for task in tasks:
        assert any(entry["text"] == "Remarks: Filters ordered, INV-04 fan noisy" for entry in task.log)

    bundled = openpyxl.load_workbook("frontend/static/templates/PM_Task_Plan_Template.xlsx")
    assert bundled["PM Tasks"].cell(1, 12).value == "Remarks"


def test_washing_setup_uses_remarks_and_still_reads_notes(client, db):
    _plant(db, "WSR01", "Setup Remarks Plant")
    login_user(client, db, role="O&M Head", name="Setup Head", email="setup-remarks@example.com")
    template = load_workbook(BytesIO(client.get("/api/scheduled/washing/template").content))
    assert [c.value for c in template["Washing setup"][1]][-1] == "Remarks"
    bundled = openpyxl.load_workbook("frontend/static/templates/Washing_Setup_Template.xlsx")
    assert [c.value for c in bundled["Washing setup"][1]][-1] == "Remarks"

    for header, value in (("Remarks", "Tanker every 3 days"), ("Notes", "South gate locked Sundays")):
        text = f"Site code,Number of tables,{header}\nWSR01,10,{value}\n"
        result = client.post("/api/scheduled/washing/import?dry_run=false", files={"file": ("s.csv", text.encode(), "text/csv")})
        assert result.status_code == 200 and result.json()["problems"] == [], result.text
        db.expire_all()
        assert db.query(WashingConfig).one().notes == value
