"""Build the Plant side and GSS side PM upload templates from the 30-day PM cycle.

    python scripts/build_pm_side_templates.py

Writes frontend/static/templates/PM_Plant_Side_Template.xlsx and PM_GSS_Side_Template.xlsx. Both use the
same columns as PM_Task_Plan_Template.xlsx except Start date: the date is chosen in the portal's upload
dialog (cycle day 1) and the order of the cycle sets every equipment's days. The workbook Keywords
(solar91-pm-side:plant / :gss) let the importer reject equipment that belongs to the other side.
Re-run after changing app/pmplan/cycle.py.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.pmplan.cycle import CYCLE, CYCLE_DAYS, MISC, phase_for  # noqa: E402

OUT = ROOT / "frontend" / "static" / "templates"
BASE = OUT / "PM_Task_Plan_Template.xlsx"
FREQUENCIES = ["Daily", "Weekly", "Monthly", "Quarterly", "Half-yearly", "Annual"]
SAFETY = ["—", "PTW required", "LOTO required", "PTW + LOTO required"]
COLUMNS = [  # (header, width); headers are matched by the importer (app/pmplan/router.py::_header_key)
    ("Equipment type", 26), ("Plant(s)", 26), ("Task - what to do", 42), ("Frequency - when", 16),
    ("Assignee(s)", 26), ("Safety needed", 22), ("Est. time (hours)", 15), ("How to do it", 42),
    ("No. of units", 13), ("Units per day", 13), ("Remarks", 34), ("Order in PM cycle (info only)", 34),
]
GREEN, GREY, YELLOW, INFO = "FF0F5C2E", "FFEEF2EF", "FFFFF9DB", "FFF2F2F2"
THIN = Side(style="thin", color="FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (equipment type, task, safety, est. hours, how to do it, no. of units)
ROWS = {
    "plant": [
        ("Inverter", "Inverter preventive maintenance", "PTW + LOTO required", 2,
         "1. Isolate. 2. Clean fans and filters. 3. Check DC / AC terminal torque. 4. Review alarms. 5. Record readings.", None),
        ("ACCB Panel", "ACCB panel PM (ACCB-1 + ACCB-2, one job)", "PTW + LOTO required", 2,
         "1. Isolate. 2. Breaker operation check. 3. Busbar / terminal thermal scan. 4. Cable lugs. 5. Panel cleaning. "
         "Do ACCB-1 and ACCB-2 in the same job.", 1),
        ("MCCB", "MCCB PM", "PTW + LOTO required", 1,
         "1. Visual / thermal inspection. 2. Terminal tightening. 3. Trip and manual operation test.", None),
        ("VCB Panel", "VCB panel PM (plant side)", "PTW + LOTO required", 3,
         "1. Breaker mechanical operation. 2. Contacts / terminals. 3. Relay & trip-circuit test. 4. Interlocks.", None),
        ("Transformer", "IDT transformer PM", "PTW + LOTO required", 3,
         "1. Oil level / leakage. 2. Bushings & breather. 3. Temperature / protection checks. 4. Earthing & terminals.", None),
        ("Aux. Transformer", "Aux. transformer PM", "PTW + LOTO required", 2,
         "1. Oil / enclosure. 2. Bushings & cables. 3. Voltage / current reading. 4. Earthing continuity.", None),
        ("4 Pole", "4 Pole structure PM (plant side)", "PTW required", 2,
         "1. Structure & foundation. 2. Insulators / jumpers. 3. Earthing connection. 4. Hardware tightening.", None),
        ("Earthing", "Earth pit maintenance", "—", 2,
         "1. Earth-pit condition. 2. Resistance measurement. 3. Connection tightening. 4. Watering / treatment.", None),
        ("Miscellaneous", "Miscellaneous PM (Street Light, ERT, …)", "—", 2,
         "1. Street Light. 2. ERT. 3. Any other small item at the plant. All in one job.", 1),
        ("Transmission line", "Transmission line patrol & PM (plant side)", "PTW required", 4,
         "1. Line patrol. 2. Conductor / sag. 3. Insulators / jumpers. 4. Poles / towers. 5. Vegetation & ROW.", None),
    ],
    "gss": [
        ("VCB - GSS side", "GSS VCB PM", "PTW + LOTO required", 3,
         "1. Breaker mechanical operation. 2. Contacts / terminals. 3. Relay & trip-circuit test. 4. Interlocks.", None),
        ("4 Pole - GSS side", "GSS 4 Pole structure PM", "PTW required", 2,
         "1. Structure & foundation. 2. Insulators / jumpers. 3. Earthing connection. 4. Hardware tightening.", None),
        ("Transmission line - GSS side", "GSS transmission line patrol & PM", "PTW required", 4,
         "1. Line patrol. 2. Conductor / sag. 3. Insulators / jumpers. 4. Poles / towers. 5. Vegetation & ROW.", None),
    ],
}
TITLE = {"plant": "Plant side", "gss": "GSS side"}


def _days(p: dict) -> str:
    return f"day {p['start_day']}" if p["days"] == 1 else f"days {p['start_day']}–{p['end_day']}"


def _order_note(etype: str) -> str:
    p = phase_for(etype)
    step = [q for q in CYCLE if q["area"] == p["area"]].index(p) + 1
    return f"{p['area']} step {step} · cycle {_days(p)}"


def _base_types_and_ideas() -> tuple[list[str], dict[str, str]]:
    wb = load_workbook(BASE, read_only=True)
    types = [r[0] for r in wb["Lists"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    ideas = {r[0]: r[1] for r in wb["Task ideas"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    ideas.setdefault("4 Pole - GSS side", ideas.get("4 Pole"))
    ideas.setdefault("Transmission line - GSS side", ideas.get("Transmission line"))
    ideas.setdefault(MISC, "Street Light, ERT and other small items — one job per plant")
    return types, ideas


def _side_types(side: str, all_types: list[str]) -> list[str]:
    if side == "gss":
        return [t for p in CYCLE if p["area"] == "GSS side" for t in p["equipment_types"]]
    named = [t for p in CYCLE if p["area"] == "Plant side" for t in p["equipment_types"]]
    return named + [t for t in all_types if t not in named and phase_for(t)["area"] != "GSS side"]


def _how_to(ws, side: str) -> None:
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 3, 26, 92
    other = "GSS side" if side == "plant" else "Plant side"
    phases = [p for p in CYCLE if p["area"] == TITLE[side]]
    def what(p: dict) -> str:
        if p["phase"] == MISC:
            return "Street Light, ERT, … as one Miscellaneous job; any other equipment type not listed here"
        return ", ".join(p["equipment_types"])
    order = "\n".join(f"{i}. {p['phase']} ({what(p)}) — cycle {_days(p)}" for i, p in enumerate(phases, 1))
    ws["B2"] = f"Solar91 O&M — {TITLE[side]} Preventive Maintenance plan"
    ws["B2"].font = Font(name="Arial", size=16, bold=True, color=GREEN)
    ws["B3"] = (f"Only {TITLE[side]} equipment goes in this file. {other} equipment is rejected on upload — "
                f"use the {other} template for it. One row = one PM job.")
    ws["B3"].font = Font(name="Arial", size=10, color="FF555555")
    rows = [
        ("Order of work", f"The portal does the {TITLE[side]} equipment in this order:\n{order}"
         + ("\nThe GSS side (VCB, 4 Pole, Transmission line) follows after the plant side is finished."
            if side == "plant" else "\nThe GSS side starts after the plant side work (Inverter … Transmission line) at the plant.")),
        ("Date", "There is no date column. Choose the cycle start date (day 1) in the portal when you upload: "
                 "Preventive Maintenance → Import CSV / Excel. Every equipment's dates follow from the order above."),
        ("No overlap", "If earlier work at a plant takes longer than its days (for example 20 inverters at 2 a day = 10 days), "
                       "the next equipment starts the day after it finishes and everything after it moves too. "
                       "The portal shows the final dates before anything is saved."),
        ("Equipment type", "Pick from the dropdown. It lists only " + TITLE[side] + " equipment."),
        ("Plant(s)", "Plant name or code exactly as in the portal; several plants separated by commas, or All plants. "
                     "Rows without a plant are saved as PM job templates only (no tasks)."),
        ("Task - what to do", "Short name of the job. Keep Plant side and GSS side task names different."),
        ("Frequency - when", "Monthly for the 30-day cycle (each unit comes back on the same cycle day next month)."),
        ("Assignee(s)", "O&M team member(s) with the Assignee role, separated by commas."),
        ("Safety needed", "— (none), PTW required, LOTO required, or PTW + LOTO required."),
        ("Est. time (hours)", "Approximate hours per unit, as a number."),
        ("How to do it", "Steps / checklist in plain words."),
        ("No. of units", "Optional. How many units the plant has. Blank = the plant's Unit configuration "
                         "or Equipment Register. ACCB Panel (ACCB-1 + ACCB-2) and Miscellaneous are always one job per plant."),
        ("Units per day", "Optional. Most units done in one day. Blank = spread the units over the equipment's cycle days."),
        ("Remarks", "Optional. Added to the log of every task the row creates."),
        ("Order in PM cycle", "For information only; the portal ignores this column."),
    ]
    ws["B5"], ws["C5"] = "Column / topic", "What to enter"
    for c in (ws["B5"], ws["C5"]):
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFFFF"), PatternFill("solid", fgColor=GREEN)
    for i, (k, text) in enumerate(rows, start=6):
        ws.cell(i, 2, k).font = Font(name="Arial", bold=True)
        ws.cell(i, 3, text).font = Font(name="Arial")
        for col in (2, 3):
            ws.cell(i, col).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(i, col).border = BOX


def build(side: str) -> Path:
    all_types, ideas = _base_types_and_ideas()
    types = _side_types(side, all_types)
    wb = Workbook()
    _how_to(wb.active, side)
    wb.active.title = "How to fill"

    ws = wb.create_sheet("PM Tasks")
    for col, (head, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(1, col, head)
        c.font, c.fill, c.border = Font(name="Arial", size=11, bold=True, color="FFFFFFFF"), PatternFill("solid", fgColor=GREEN), BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    last = 200
    for r in range(2, last + 1):
        for col in range(1, len(COLUMNS) + 1):
            c = ws.cell(r, col)
            info = col == len(COLUMNS)
            c.fill, c.border = PatternFill("solid", fgColor=INFO if info else YELLOW), BOX
            c.font = Font(name="Arial", size=10, italic=info, color="FF666666" if info else "FF000000")
            c.alignment = Alignment(vertical="top", wrap_text=col in (3, 8, 11, 12))
    for r, (etype, task, safety, hours, how, units) in enumerate(ROWS[side], start=2):
        for col, value in enumerate([etype, None, task, "Monthly", None, safety, hours, how, units, None, None,
                                     _order_note(etype)], start=1):
            ws.cell(r, col, value)

    lists = wb.create_sheet("Lists")
    for col, (head, values) in enumerate([("Equipment types", types), ("Frequency", FREQUENCIES), ("Safety", SAFETY)], start=1):
        lists.cell(1, col, head).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            lists.cell(r, col, v)
    lists.column_dimensions["A"].width = 30
    for sqref, formula in ((f"A2:A{last}", f"=Lists!$A$2:$A${len(types) + 1}"),
                           (f"D2:D{last}", f"=Lists!$B$2:$B${len(FREQUENCIES) + 1}"),
                           (f"F2:F{last}", f"=Lists!$C$2:$C${len(SAFETY) + 1}")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Choose a value from the dropdown.")
        dv.add(sqref)
        ws.add_data_validation(dv)
    for sqref, kind, low in ((f"G2:G{last}", "decimal", "0"), (f"I2:J{last}", "whole", "1")):
        dv = DataValidation(type=kind, operator="greaterThanOrEqual", formula1=low, allow_blank=True)
        dv.add(sqref)
        ws.add_data_validation(dv)

    tips = wb.create_sheet("Task ideas")
    tips.append(["Equipment type", "Common checks"])
    for c in tips[1]:
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFFFF"), PatternFill("solid", fgColor=GREEN)
    for t in types:
        if ideas.get(t):
            tips.append([t, ideas[t]])
    tips.column_dimensions["A"].width, tips.column_dimensions["B"].width = 28, 95
    tips.freeze_panes = "A2"

    wb.active = 0
    wb.properties.title = f"Solar91 PM plan — {TITLE[side]}"
    wb.properties.keywords = f"solar91-pm-side:{side}"
    wb.properties.creator = "Solar91 O&M"
    path = OUT / f"PM_{'Plant' if side == 'plant' else 'GSS'}_Side_Template.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    assert CYCLE[-1]["end_day"] == CYCLE_DAYS
    for s in ("plant", "gss"):
        print(build(s).relative_to(ROOT))
