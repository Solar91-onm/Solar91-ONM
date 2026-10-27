from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "frontend" / "static" / "templates" / "Washing_Setup_Template.xlsx"
HEADERS = [
    "Site code",
    "Plant name",
    "Number of tables",
    "Washing assignee email",
    "Compliance target %",
    "Active",
    "Remarks",
]


def build_template(output: Path = OUTPUT) -> None:
    workbook = Workbook()
    how_to = workbook.active
    how_to.title = "How to fill"
    for line in (
        "Use one row for each plant.",
        "Site code must match a plant already in Solar91.",
        "For a new setup, enter the number of tables. Assignee and other settings can be added later.",
        "Compliance target % is optional and must be from 1 to 100; the default is 85. Each washing cycle (days 1-10, 11-20, 21-month end) is compared with it.",
        "For updates, blank cells mean leave the saved value as it is.",
        "Remarks: the current situation at the plant (e.g. water tanker every 3 days, south gate locked on Sundays).",
    ):
        how_to.append([line])
    how_to.column_dimensions["A"].width = 95
    how_to["A1"].font = Font(bold=True, color="FFFFFF")
    how_to["A1"].fill = PatternFill("solid", fgColor="176B3A")
    how_to.freeze_panes = "A2"

    setup = workbook.create_sheet("Washing setup")
    setup.append(HEADERS)
    setup.freeze_panes = "A2"
    setup.auto_filter.ref = "A1:G1"
    widths = (18, 32, 20, 34, 18, 14, 50)
    for index, width in enumerate(widths, start=1):
        setup.column_dimensions[setup.cell(1, index).column_letter].width = width
    for cell in setup[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B3A")
        cell.alignment = Alignment(wrap_text=True)

    integer_tables = DataValidation(type="whole", operator="between", formula1="1", formula2="5000")
    integer_target = DataValidation(type="whole", operator="between", formula1="1", formula2="100")
    active = DataValidation(type="list", formula1='"Yes,No"')
    for validation in (integer_tables, integer_target, active):
        setup.add_data_validation(validation)
    integer_tables.add("C2:C501")
    integer_target.add("E2:E501")
    active.add("F2:F501")

    lists = workbook.create_sheet("Lists")
    lists.append(["Assignee emails", "Site codes"])
    lists.column_dimensions["A"].width = 36
    lists.column_dimensions["B"].width = 24
    lists.sheet_state = "hidden"
    assignees = DataValidation(type="list", formula1="'Lists'!$A$2:$A$501", allow_blank=True)
    setup.add_data_validation(assignees)
    assignees.add("D2:D501")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


if __name__ == "__main__":
    build_template()
